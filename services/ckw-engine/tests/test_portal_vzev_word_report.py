from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from ckw_script.portal_analysis import create_project, ingest_portal_file, set_pv_battery_consumers
from ckw_script.portal_api import PortalSettings, create_app
from ckw_script.portal_vzev_word_report import create_vzev_word_report


def _plant_cache(tmp_path) -> str:
    path = tmp_path / "plants.csv"
    path.write_text(
        "\n".join(
            [
                "Gemeinde;Strasse;Hausnummer;PLZ;Energietraeger;Installierte Leistung",
                "Sursee;Bahnhofstrasse;1;6210;Photovoltaik;30",
            ]
        ),
        encoding="utf-8",
    )
    return str(path)


def _ckw_csv(*, contract: str, column: str, values: list[float]) -> bytes:
    rows = [
        "Kunde: Test AG",
        "Kundennummer: 1",
        "Adresse: Bahnhofstrasse 1, 6210 Sursee",
        f"Vertrag: {contract}",
        "Zaehlpunkt: CH1003601234500000000000000000001",
        ";;",
        "Datum der Auswertung:01.01.2026",
        "Auswertungszeitraum:01.01.2025 - 01.01.2025",
        "Einheit:Viertelstunde",
        "",
        f"Zeitraum;{column}",
    ]
    for index, value in enumerate(values):
        rows.append(f"01.01.2025 07:{index * 15:02d};{value}")
    return "\n".join(rows).encode("utf-8")


def _project_with_simulation_inputs(tmp_path):
    db_path = tmp_path / "portal.duckdb"
    project = create_project(name="vZEV Word Test", municipality="Sursee", db_path=db_path)
    solar = ingest_portal_file(
        project_id=project["project_id"],
        original_filename="solar.csv",
        content=_ckw_csv(contract="PVA", column="Produktion (kWh)", values=[1.0, 4.0, 0.0, 2.0]),
        upload_kind="solar",
        db_path=db_path,
    )
    consumption = ingest_portal_file(
        project_id=project["project_id"],
        original_filename="verbrauch.csv",
        content=_ckw_csv(contract="Haus A", column="Energieverbrauch (kWh)", values=[3.0, 2.0, 1.0, 1.0]),
        upload_kind="consumption",
        db_path=db_path,
    )
    set_pv_battery_consumers(
        project_id=project["project_id"],
        solar_file_id=solar.file_id,
        consumer_file_ids=[consumption.file_id],
        tariff_year=2025,
        db_path=db_path,
    )
    return db_path, project, consumption


def test_vzev_word_report_contains_complete_workspace(tmp_path) -> None:
    db_path, project, consumption = _project_with_simulation_inputs(tmp_path)

    report = create_vzev_word_report(
        project_id=project["project_id"],
        db_path=db_path,
        plant_cache_path=_plant_cache(tmp_path),
        tariff_year=2025,
        consumption_file_id=consumption.file_id,
        load_shifting_year=2025,
        fixed_local_energy_rp_per_kwh=Decimal("5.00"),
    )

    document = Document(BytesIO(report.content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert report.filename == "vzev-simulationsbericht-vzev-word-test.docx"
    assert paragraphs[0] == "vZEV-Simulationsbericht"
    for heading in (
        "Eingabedaten",
        "Solarenergie",
        "Gemessener vZEV-Energiefluss",
        "Lastverschiebung Nacht zu Tag",
        "vZEV-Simulation",
        "Teilnehmer – Energie und Preise",
        "Teilnehmer – Kostenvergleich",
        "Tarif- und Leistungspreisänderung je Messpunkt",
        "Kostenänderung je Messpunkt",
        "Batterieplatzierung",
        "Berechnungen und Annahmen",
    ):
        assert heading in paragraphs
    assert "solar.csv" in table_text
    assert "verbrauch.csv" in table_text
    assert "Haus A" in table_text
    assert "CH1003601234500000000000000000001" in table_text
    assert "Änderung Leistung" in table_text
    assert "gegen die jeweilige 80%-Obergrenze geprüft" in "\n".join(paragraphs)
    assert "Plus bedeutet Mehrkosten" in "\n".join(paragraphs)
    assert "PV-Tarif: fixer Tarif von 5.00 Rp./kWh" in "\n".join(paragraphs)
    assert len(document.inline_shapes) >= 5


def test_vzev_word_report_keeps_overridden_tariff_and_exports_warning(tmp_path) -> None:
    db_path, project, consumption = _project_with_simulation_inputs(tmp_path)

    report = create_vzev_word_report(
        project_id=project["project_id"],
        db_path=db_path,
        plant_cache_path=_plant_cache(tmp_path),
        tariff_year=2025,
        consumption_file_id=consumption.file_id,
        fixed_local_energy_rp_per_kwh=Decimal("1000.00"),
    )

    document = Document(BytesIO(report.content))
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert "PV-Tarif: fixer Tarif von 1'000.00 Rp./kWh" in paragraph_text
    assert "Die Simulation verwendet den eingegebenen Tarif trotzdem unverändert" in paragraph_text
    assert "Über 80%" in table_text


def test_vzev_word_report_endpoint_and_incomplete_project(tmp_path) -> None:
    db_path = tmp_path / "portal.duckdb"
    project = create_project(name="Noch leer", municipality="Sursee", db_path=db_path)
    app = create_app(
        PortalSettings(
            portal_db_path=str(db_path),
            pv_plant_cache_path=_plant_cache(tmp_path),
            frontend_dist_path=str(tmp_path / "missing-frontend"),
        )
    )

    response = TestClient(app).get(f"/api/projects/{project['project_id']}/vzev-word-report")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert response.headers["content-disposition"] == (
        'attachment; filename="vzev-simulationsbericht-noch-leer.docx"'
    )
    document = Document(BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Noch keine Messdatei importiert." in text
    assert "Simulation nicht verfügbar" in text
