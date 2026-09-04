from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from ckw_script.portal_analysis import create_project, ingest_portal_file
from ckw_script.portal_api import PortalSettings, create_app
from ckw_script.portal_excel_export import EXCEL_MEDIA_TYPE, create_project_excel_export


def _column_by_header(sheet, header: str) -> int:
    for column in range(1, sheet.max_column + 1):
        if sheet.cell(row=1, column=column).value == header:
            return column
    raise AssertionError(f"Header not found: {header}")


def _ckw_csv(*, contract: str, column: str, values: list[float]) -> bytes:
    rows = [
        "Kunde: Test AG",
        "Kundennummer: 1",
        "Adresse: Bahnhofstrasse 1, 6210 Sursee",
        f"Vertrag: {contract}",
        "Zaehlpunkt: CH1003601234500000000000000000001",
        ";;",
        "Datum der Auswertung:01.01.2026",
        "Auswertungszeitraum:01.01.2025 - 01.01.2025",
        "Einheit:Viertelstunde",
        "",
        f"Zeitraum;{column}",
    ]
    for index, value in enumerate(values):
        rows.append(f"01.01.2025 00:{index * 15:02d};{value}")
    return "\n".join(rows).encode("utf-8")


def test_create_project_excel_export_contains_validation_sheets_and_formulas(tmp_path) -> None:
    db_path = tmp_path / "portal.duckdb"
    project = create_project(name="Excel Export", municipality="Sursee", db_path=db_path)
    ingest_portal_file(
        project_id=project["project_id"],
        original_filename="solar.csv",
        content=_ckw_csv(contract="Ueberschuss PVA", column="Rücklieferung (kWh)", values=[1.0, 2.0]),
        upload_kind="solar",
        db_path=db_path,
    )
    ingest_portal_file(
        project_id=project["project_id"],
        original_filename="load.csv",
        content=_ckw_csv(contract="Allgemein", column="Energieverbrauch (kWh)", values=[0.5, 3.0]),
        upload_kind="consumption",
        db_path=db_path,
    )

    export = create_project_excel_export(project_id=project["project_id"], tariff_year=2025, db_path=db_path)
    workbook = load_workbook(BytesIO(export.content), data_only=False)

    assert export.filename == "leg-validierung-excel-export.xlsx"
    assert workbook.sheetnames == [
        "Projekt",
        "Dateien",
        "Messreihen 15min",
        "LEG Rechnung 15min",
        "Monate Summen",
        "Optimierung",
        "Annahmen Warnungen",
    ]

    raw_sheet = workbook["Messreihen 15min"]
    assert raw_sheet.max_row == 35_041
    assert raw_sheet["A1"].value == "timestamp"
    assert raw_sheet["A2"].value.isoformat() == "2025-01-01T00:00:00"
    solar_column = _column_by_header(raw_sheet, "Ueberschuss PVA | PV-Ueberschuss kWh")
    load_column = _column_by_header(raw_sheet, "Allgemein | Verbrauch kWh")
    assert raw_sheet.cell(row=2, column=solar_column).value == pytest.approx(1.0)
    assert raw_sheet.cell(row=3, column=load_column).value == pytest.approx(3.0)

    leg_sheet = workbook["LEG Rechnung 15min"]
    assert leg_sheet.max_row == 35_041
    assert leg_sheet["B2"].value == pytest.approx(0.5)
    assert leg_sheet["C2"].value == pytest.approx(1.0)
    assert leg_sheet["D2"].value == "=MIN(B2,C2)"
    assert leg_sheet["E2"].value == "=MAX(C2-D2,0)"
    assert leg_sheet["J2"].value == pytest.approx(0.4)
    assert leg_sheet["K2"].value == "=D2*H2/100"
    assert leg_sheet["L2"].value == "=D2*(G2-H2)/100"
    assert leg_sheet["M2"].value == "=D2*I2*J2/100"
    assert leg_sheet["N2"].value == "=L2+M2"

    files_sheet = workbook["Dateien"]
    assert any(
        files_sheet.cell(row=row, column=1).value == "Ueberschuss PVA | PV-Ueberschuss kWh"
        for row in range(1, files_sheet.max_row + 1)
    )
    assert any(
        files_sheet.cell(row=row, column=1).value == "Allgemein | Verbrauch kWh"
        for row in range(1, files_sheet.max_row + 1)
    )

    assumptions = [
        workbook["Annahmen Warnungen"].cell(row=row, column=2).value
        for row in range(2, workbook["Annahmen Warnungen"].max_row + 1)
    ]
    assert any("Excel-Formeln pro Intervall" in str(value) for value in assumptions)


def test_excel_export_endpoint_downloads_xlsx(tmp_path) -> None:
    db_path = tmp_path / "portal.duckdb"
    project = create_project(name="API Excel", municipality="Sursee", db_path=db_path)
    app = create_app(
        PortalSettings(
            portal_db_path=str(db_path),
            frontend_dist_path=str(tmp_path / "missing-frontend"),
        )
    )

    response = TestClient(app).get(f"/api/projects/{project['project_id']}/excel-export")

    assert response.status_code == 200
    assert response.headers["content-type"] == EXCEL_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="leg-validierung-api-excel.xlsx"'
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames[0] == "Projekt"
