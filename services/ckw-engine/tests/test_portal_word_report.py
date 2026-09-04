from __future__ import annotations

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from ckw_script.portal_analysis import create_project, ingest_portal_file
from ckw_script.portal_api import PortalSettings, create_app
from ckw_script.portal_word_report import create_project_word_report


def _table_by_header(document: Document, header: str):
    return next(table for table in document.tables if table.cell(0, 0).text == header)


def _plant_cache(tmp_path) -> str:
    cache_path = tmp_path / "plants.csv"
    cache_path.write_text(
        "\n".join(
            [
                "Gemeinde;Strasse;Hausnummer;PLZ;Energietraeger;Installierte Leistung",
                "Sursee;Bahnhofstrasse;1;6210;Photovoltaik;30",
            ]
        ),
        encoding="utf-8",
    )
    return str(cache_path)


def _ckw_csv(*, contract: str, column: str, values: list[float]) -> bytes:
    rows = [
        "Kunde: Test AG",
        "Kundennummer: 1",
        "Adresse: Bahnhofstrasse 1, 6210 Sursee",
        f"Vertrag: {contract}",
        "Zaehlpunkt: CH1003601234500000000000000000001",
        ";;",
        "Datum der Auswertung:01.01.2026",
        "Auswertungszeitraum:01.01.2025 - 01.01.2025",
        "Einheit:Viertelstunde",
        "",
        f"Zeitraum;{column}",
    ]
    for index, value in enumerate(values):
        rows.append(f"01.01.2025 00:{index * 15:02d};{value}")
    return "\n".join(rows).encode("utf-8")


def test_create_project_word_report_is_editable_docx(tmp_path) -> None:
    db_path = tmp_path / "portal.duckdb"
    project = create_project(name="Sursee Test", municipality="Sursee", db_path=db_path)
    ingest_portal_file(
        project_id=project["project_id"],
        original_filename="solar.csv",
        content=_ckw_csv(contract="Ueberschuss PVA", column="Rücklieferung (kWh)", values=[1.0, 2.0]),
        upload_kind="solar",
        db_path=db_path,
    )
    ingest_portal_file(
        project_id=project["project_id"],
        original_filename="load.csv",
        content=_ckw_csv(contract="Allgemein", column="Energieverbrauch (kWh)", values=[0.5, 3.0]),
        upload_kind="consumption",
        db_path=db_path,
    )

    report = create_project_word_report(
        project_id=project["project_id"],
        db_path=db_path,
        plant_cache_path=_plant_cache(tmp_path),
    )
    document = Document(BytesIO(report.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    paragraph_texts = [paragraph.text for paragraph in document.paragraphs]

    assert report.filename == "leg-potenzial-sursee-test.docx"
    assert text.startswith("LEG-Potenzialbericht")
    assert "Projekt: Sursee Test" in text
    assert "Lokal verbrauchter Solarstrom: 2 kWh" in text
    assert "Methode" in paragraph_texts
    assert "Produktionseinheiten" in paragraph_texts
    assert "Verbleibender Solarüberschuss" in paragraph_texts
    assert "Verteilung nach Monat" in paragraph_texts
    assert "Verteilung nach Datum und Tageszeit" in paragraph_texts
    assert "Höchste 15-Minuten-Intervalle" in paragraph_texts
    vzev_index = paragraph_texts.index("VZEV-Potenzial")
    assert paragraph_texts[vzev_index + 1] == "Monatliche Energieübersicht"
    plant_table = _table_by_header(document, "Adresse")
    assert plant_table.cell(1, 0).text == "Bahnhofstrasse 1"
    assert plant_table.cell(1, 1).text == "Sursee"
    assert plant_table.cell(1, 2).text == "30.0 kWp"
    monthly_table = _table_by_header(document, "Monat")
    assert monthly_table.cell(1, 0).text == "01.2025"
    assert monthly_table.cell(1, 1).text == "3 kWh"
    assert monthly_table.cell(1, 2).text == "4 kWh"
    assert monthly_table.cell(1, 3).text == "2 kWh"
    assert _table_by_header(document, "Uhrzeit").cell(1, 0).text == "00:00"
    assert _table_by_header(document, "Zeitpunkt").cell(1, 3).text == "0 kWh"
    assert _table_by_header(document, "Datei").cell(1, 0).text == "solar.csv"
    assert _table_by_header(document, "Teilnehmer").cell(1, 0).text == "Allgemein"
    assert len(document.inline_shapes) == 4


def test_word_report_endpoint_downloads_docx(tmp_path) -> None:
    db_path = tmp_path / "portal.duckdb"
    project = create_project(name="API Export", municipality="Sursee", db_path=db_path)
    app = create_app(
        PortalSettings(
            portal_db_path=str(db_path),
            pv_plant_cache_path=_plant_cache(tmp_path),
            frontend_dist_path=str(tmp_path / "missing-frontend"),
        )
    )

    response = TestClient(app).get(f"/api/projects/{project['project_id']}/word-report")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert response.headers["content-disposition"] == 'attachment; filename="leg-potenzial-api-export.docx"'
    assert Document(BytesIO(response.content)).paragraphs[0].text == "LEG-Potenzialbericht"
