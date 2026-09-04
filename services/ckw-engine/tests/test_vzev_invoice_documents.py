from __future__ import annotations

import json
import os
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
import pytest

from ckw_script.portal_api import PortalSettings, create_app
from ckw_script.vzev_invoice_bundle import InvoiceSource, create_invoice_batch, create_invoice_bundle


def _private_vzev_fixture_directory() -> Path:
    configured = os.environ.get("CKW_PRIVATE_FIXTURES_DIR")
    directory = Path(configured) if configured else Path(__file__).resolve().parents[1] / "data/vzev"
    if not directory.exists():
        pytest.skip("private CKW reconciliation fixtures are not present in this code-only checkout")
    return directory


def _monthly_vzev_csv(month: int = 3) -> bytes:
    timestamp = f"01.{month:02d}.2026"
    rows = [
        "Zählpunkt Webbezeichnung;PV: Überschuss PVA;;virtuelles AO vZEV: vZEV Summe Bezug;;virtuelles AO vZEV: vZEV Summe Überschuss;;Haus A: Allgemein;;Haus B: Allgemein;",
        "Metering-Code;CH1003601234500000000000000000001;;CH1003601234500000000000000000002;;CH1003601234500000000000000000003;;CH1003601234500000000000000000004;;CH1003601234500000000000000000005;",
        "Linie Webbezeichnung;Wirkenergie Total - (Bezug);;Wirkenergie Total + (Abgabe);;Wirkenergie Total - (Bezug);;Wirkenergie Total + (Abgabe);;Wirkenergie Total + (Abgabe);",
        "Obis-Code;1-5:2.29.0*255;;1-5:1.29.0*255;;1-5:2.29.0*255;;1-5:1.29.0*255;;1-5:1.29.0*255;",
        "Einheit;kWh;;kWh;;kWh;;kWh;;kWh;",
        ";0;;0;;0;;0;;0;",
        f"{timestamp} 00:15;1.0;W;2.0;W;0;W;1.0;W;2.0;W",
        f"{timestamp} 00:30;1.0;W;4.0;W;0;W;2.0;W;3.0;W",
    ]
    return "\n".join(rows).encode("cp1252")


def _document_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def test_invoice_bundle_contains_one_editable_word_file_per_consumer() -> None:
    bundle = create_invoice_bundle(
        csv_content=_monthly_vzev_csv(),
        original_filename="monthly.csv",
        issuer="vZEV Test",
        invoice_date="2026-04-01",
        payment_terms="   ",
    )

    with ZipFile(BytesIO(bundle.content)) as archive:
        names = archive.namelist()
        docx_names = [name for name in names if name.endswith(".docx")]
        participant_docx_names = [name for name in docx_names if "gemeinschaftsuebersicht" not in name]
        overview_name = next(name for name in docx_names if "gemeinschaftsuebersicht" in name)
        assert len(participant_docx_names) == 2
        assert len(docx_names) == 3
        assert any(name.endswith("_index.json") for name in names)
        participant_text = _document_text(archive.read(participant_docx_names[0])) + _document_text(
            archive.read(participant_docx_names[1])
        )
        first_invoice = Document(BytesIO(archive.read(participant_docx_names[0])))
        paragraph_texts = [paragraph.text for paragraph in first_invoice.paragraphs]
        issuer_index = paragraph_texts.index("Rechnungssteller: vZEV Test")
        invoice_date_index = paragraph_texts.index("Rechnungsdatum: 2026-04-01")
        participant_index = next(
            index for index, text in enumerate(paragraph_texts) if text.startswith("Teilnehmer: ")
        )
        positions_index = paragraph_texts.index("Rechnungspositionen")
        total_index = paragraph_texts.index("Total")
        details_index = paragraph_texts.index("Rechnungsdetails")
        assert issuer_index < invoice_date_index < participant_index < positions_index
        assert positions_index < total_index < details_index
        assert first_invoice.paragraphs[details_index - 1]._p.xpath(
            ".//w:br[@w:type='page']"
        )
        details_table = next(
            table
            for table in first_invoice.tables
            if table.rows and table.rows[0].cells[0].text == "Teilnehmer"
        )
        assert [cell.text for cell in details_table.rows[0].cells] != ["Feld", "Wert"]
        assert details_table.autofit is False
        assert details_table.columns[1].width > details_table.columns[0].width
        assert "Haus A: Allgemein" in participant_text
        assert "Vergleich mit Bezug ohne vZEV" not in participant_text
        assert "Zahlungsfrist" not in participant_text
        assert "Berechnungsmethode" not in participant_text
        assert "Hinweise" not in participant_text
        assert "80% × vollständiger CKW-Standardproduktpreis" in participant_text
        assert "Variable Netznutzung inkl. Systemdienste" in participant_text
        assert "Grund- und Messpreis" in participant_text
        assert "Vollständiger Standardproduktpreis" in participant_text
        overview_text = _document_text(archive.read(overview_name))
        assert "vZEV Gemeinschaftsübersicht" in overview_text
        assert "Zahlbetrag an CKW" in overview_text
        assert "Energiemengen" in overview_text
        assert "Kostenaufteilung" in overview_text
        assert "PV-Energiemengen" in overview_text
        assert "PV-Auszahlungen" in overview_text
        assert "Export gemessen" in overview_text
        assert "Export abgerechnet" in overview_text

        index_name = next(name for name in names if name.endswith("_index.json"))
        index = json.loads(archive.read(index_name))
        assert index["source_file"] == "monthly.csv"
        assert index["invoice_count"] == 2
        assert index["document_format"] == "docx"
        assert index["community_overview"].endswith("gemeinschaftsuebersicht.docx")
        for invoice in index["invoices"]:
            assert invoice["consumption_kwh"] > 0
            assert round(invoice["pv_payment_chf"] + invoice["ckw_payment_chf"], 2) == invoice["total_chf"]


def test_invoice_api_returns_downloadable_zip(tmp_path) -> None:
    app = create_app(
        PortalSettings(
            portal_db_path=str(tmp_path / "portal.duckdb"),
            pv_plant_cache_path=str(tmp_path / "plants.csv"),
            frontend_dist_path=str(tmp_path / "missing-frontend"),
        )
    )
    response = TestClient(app).post(
        "/api/vzev/invoices",
        files={"csv_file": ("monthly.csv", _monthly_vzev_csv(), "text/csv")},
        data={"issuer": "vZEV Test", "invoice_date": "2026-04-01"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert response.headers["x-invoice-count"] == "2"
    assert "vzev-rechnungen-202603.zip" in response.headers["content-disposition"]
    with ZipFile(BytesIO(response.content)) as archive:
        assert len([name for name in archive.namelist() if name.endswith(".docx")]) == 3


def test_invoice_batch_packages_months_chronologically_with_a_batch_index() -> None:
    bundle = create_invoice_batch(
        csv_files=[
            InvoiceSource(content=_monthly_vzev_csv(4), original_filename="april.csv"),
            InvoiceSource(content=_monthly_vzev_csv(3), original_filename="march.csv"),
        ],
        issuer="vZEV Test",
        invoice_date="2026-05-01",
    )

    assert bundle.period_keys == ("202603", "202604")
    assert bundle.period_count == 2
    assert bundle.invoice_count == 4
    assert bundle.filename == "vzev-rechnungen-202603-bis-202604.zip"
    with ZipFile(BytesIO(bundle.content)) as archive:
        names = archive.namelist()
        assert "_batch_index.json" in names
        assert any(name.startswith("vzev-202603/") for name in names)
        assert any(name.startswith("vzev-202604/") for name in names)
        batch_index = json.loads(archive.read("_batch_index.json"))
        assert batch_index["period_count"] == 2
        assert batch_index["invoice_count"] == 4
        assert [period["source_file"] for period in batch_index["periods"]] == [
            "march.csv",
            "april.csv",
        ]
        overview_name = batch_index["multi_month_overview"]
        assert overview_name in names
        overview_text = _document_text(archive.read(overview_name))
        assert "vZEV Mehrmonatsübersicht" in overview_text
        assert "Geld an PV" in overview_text
        assert "Geld an CKW" in overview_text
        assert "März 2026" in overview_text
        assert "April 2026" in overview_text
        haus_a = next(
            participant
            for participant in batch_index["participants"]
            if participant["participant"] == "Haus A: Allgemein"
        )
        assert len(haus_a["months"]) == 2
        assert haus_a["totals"]["consumption_kwh"] == 6.0
        for month in haus_a["months"]:
            assert round(month["pv_payment_chf"] + month["ckw_payment_chf"], 2) == month["total_chf"]


def test_invoice_api_accepts_multiple_monthly_files(tmp_path) -> None:
    app = create_app(
        PortalSettings(
            portal_db_path=str(tmp_path / "portal.duckdb"),
            pv_plant_cache_path=str(tmp_path / "plants.csv"),
            frontend_dist_path=str(tmp_path / "missing-frontend"),
        )
    )
    response = TestClient(app).post(
        "/api/vzev/invoices",
        files=[
            ("csv_file", ("march.csv", _monthly_vzev_csv(3), "text/csv")),
            ("csv_file", ("april.csv", _monthly_vzev_csv(4), "text/csv")),
        ],
        data={"issuer": "vZEV Test", "invoice_date": "2026-05-01", "payment_terms": ""},
    )

    assert response.status_code == 200
    assert response.headers["x-period-count"] == "2"
    assert response.headers["x-invoice-count"] == "4"
    assert "vzev-rechnungen-202603-bis-202604.zip" in response.headers["content-disposition"]
    with ZipFile(BytesIO(response.content)) as archive:
        participant_names = [
            name
            for name in archive.namelist()
            if name.endswith(".docx") and "uebersicht" not in name
        ]
        assert len(participant_names) == 4
        for participant_name in participant_names:
            text = _document_text(archive.read(participant_name))
            assert "Rechnungsdatum: 2026-05-01" in text
            assert "Zahlungsfrist" not in text


def test_invoice_batch_rejects_duplicate_periods() -> None:
    with pytest.raises(ValueError, match="Abrechnungsperiode 202603 wurde bereits"):
        create_invoice_batch(
            csv_files=[
                InvoiceSource(content=_monthly_vzev_csv(3), original_filename="march-a.csv"),
                InvoiceSource(content=_monthly_vzev_csv(3), original_filename="march-b.csv"),
            ]
        )


def test_invoice_api_rejects_the_whole_batch_when_one_file_is_invalid(tmp_path) -> None:
    app = create_app(
        PortalSettings(
            portal_db_path=str(tmp_path / "portal.duckdb"),
            pv_plant_cache_path=str(tmp_path / "plants.csv"),
            frontend_dist_path=str(tmp_path / "missing-frontend"),
        )
    )
    response = TestClient(app).post(
        "/api/vzev/invoices",
        files=[
            ("csv_file", ("march.csv", _monthly_vzev_csv(3), "text/csv")),
            ("csv_file", ("broken.csv", b"not-a-ckw-export", "text/csv")),
        ],
    )

    assert response.status_code == 400
    assert "broken.csv" in response.json()["detail"]
    assert response.headers["content-type"].startswith("application/json")


def test_march_sample_reconciles_to_official_ckw_net_total() -> None:
    source = (
        _private_vzev_fixture_directory()
        / "vZEV_Allmendstrasse_4_Sursee_Einzelmesswerte_monatlich_plausibel_01.03.2026_01.04.2026.csv"
    )
    bundle = create_invoice_bundle(
        csv_content=source.read_bytes(),
        original_filename=source.name,
        issuer="vZEV Allmendstrasse 4 Sursee",
        invoice_date="2026-04-01",
    )

    with ZipFile(BytesIO(bundle.content)) as archive:
        index_name = next(name for name in archive.namelist() if name.endswith("_index.json"))
        index = json.loads(archive.read(index_name))
        assert index["gross_ckw_invoice_chf"] == 3262.85
        assert index["feed_in_compensation"]["total_chf"] == 219.71
        assert index["feed_in_compensation"]["payment_rounding_chf"] == 0.01
        assert index["net_ckw_payable_chf"] == 3043.15
        participant_total = round(sum(invoice["total_chf"] for invoice in index["invoices"]), 2)
        assert index["local_pv_energy_chf"] == 250.63
        assert participant_total == 3513.49
        assert participant_total == round(
            index["net_ckw_payable_chf"]
            + index["feed_in_compensation"]["total_chf"]
            + index["local_pv_energy_chf"],
            2,
        )


def test_real_march_to_july_batch_preserves_monthly_reconciliation() -> None:
    source_directory = _private_vzev_fixture_directory()
    source_paths = [
        source_directory
        / (
            "vZEV_Allmendstrasse_4_Sursee_Einzelmesswerte_monatlich_plausibel_"
            f"01.{month:02d}.2026_01.{month + 1:02d}.2026.csv"
        )
        for month in range(3, 8)
    ]
    sources = [
        InvoiceSource(content=path.read_bytes(), original_filename=path.name)
        for path in source_paths
    ]

    bundle = create_invoice_batch(
        csv_files=sources,
        issuer="vZEV Allmendstrasse 4 Sursee",
        invoice_date="2026-08-01",
    )

    assert bundle.period_keys == ("202603", "202604", "202605", "202606", "202607")
    assert bundle.period_count == 5
    assert bundle.invoice_count == 15
    with ZipFile(BytesIO(bundle.content)) as archive:
        batch_index = json.loads(archive.read("_batch_index.json"))
        assert batch_index["period_count"] == 5
        assert batch_index["invoice_count"] == 15
        assert len(batch_index["participants"]) == 3
        assert all(len(participant["months"]) == 5 for participant in batch_index["participants"])
        assert batch_index["multi_month_overview"] in archive.namelist()
        march_index = json.loads(archive.read("vzev-202603/_index.json"))
        assert march_index["net_ckw_payable_chf"] == 3043.15
