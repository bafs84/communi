from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import math
from pathlib import Path
import re
import textwrap
from typing import Any
import unicodedata

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from docxtpl import DocxTemplate
from PIL import Image, ImageDraw, ImageFont

from .portal_analysis import (
    DEFAULT_PLANT_CACHE,
    analyze_residual_solar,
    calculate_leg_potential,
    match_project_plants,
    optimize_leg_combinations,
    project_timeseries,
    summarize_project,
)


@dataclass(frozen=True)
class WordReport:
    filename: str
    content: bytes


def _format_number(value: object, digits: int = 0) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        number = 0.0
    return f"{number:,.{digits}f}".replace(",", "'")


def _format_kwh(value: object) -> str:
    return f"{_format_number(value)} kWh"


def _format_heatmap_kwh(value: object) -> str:
    return _format_kwh(value)


def _format_chf(value: object) -> str:
    return f"CHF {_format_number(value)}"


def _format_pct(value: object) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        number = 0.0
    return f"{number * 100:.1f}%"


def _month_label(value: object) -> str:
    try:
        return datetime.fromisoformat(str(value)).strftime("%m.%Y")
    except ValueError:
        return str(value)


def _timestamp_label(value: object) -> str:
    try:
        return datetime.fromisoformat(str(value)).strftime("%d.%m.%Y %H:%M")
    except ValueError:
        return str(value)


def _slug(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "leg-report"))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text or "leg-report"


def _default_template() -> bytes:
    document = Document()
    document.add_heading("LEG-Potenzialbericht", level=0)
    document.add_paragraph("Projekt: {{ project_name }}")
    document.add_paragraph("Gemeinde: {{ municipality }}")
    document.add_paragraph("Berichtsdatum: {{ report_date }}")
    document.add_heading("Zusammenfassung", level=1)
    document.add_paragraph("Verbrauch: {{ consumption_kwh }}")
    document.add_paragraph("Verfügbarer Solarstrom: {{ available_solar_kwh }}")
    document.add_paragraph("Lokal verbrauchter Solarstrom: {{ local_solar_kwh }}")
    document.add_paragraph("Verbleibender Netzbezug: {{ remaining_grid_import_kwh }}")
    document.add_paragraph("Vereinbarter Produzentenerlös: {{ producer_revenue_chf }}")
    document.add_paragraph("Energiepreisvorteil Verbraucher: {{ energy_price_savings_chf }}")
    document.add_paragraph("LEG-Netzkostenrabatt: {{ network_discount_chf }}")
    document.add_paragraph("Verbrauchervorteil total: {{ total_monetary_benefit_chf }}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _render_template(*, context: dict[str, str], template_path: str | Path | None) -> DocumentType:
    source: str | BytesIO
    if template_path:
        path = Path(template_path)
        if not path.is_file():
            raise ValueError(f"Word-Berichtsvorlage wurde nicht gefunden: {path}")
        source = str(path)
    else:
        source = BytesIO(_default_template())
    template = DocxTemplate(source)
    template.render(context)
    output = BytesIO()
    template.save(output)
    output.seek(0)
    return Document(output)


def _add_table(document: DocumentType, *, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    filename = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    try:
        return ImageFont.truetype(filename, size=size)
    except OSError:
        return ImageFont.load_default()


def _png_bytes(image: Image.Image) -> BytesIO:
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def _monthly_energy_chart(rows: list[dict[str, Any]]) -> BytesIO | None:
    if not rows:
        return None

    width, height = 1500, 760
    left, right, top, bottom = 120, 40, 120, 125
    plot_width = width - left - right
    plot_height = height - top - bottom
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(34, bold=True)
    label_font = _font(22)
    small_font = _font(19)
    legend_font = _font(21)
    series = [
        ("pv_excess_kwh", "PV-Überschuss", "#d79c24"),
        ("consumption_kwh", "Bedarf", "#5167a8"),
        ("local_solar_kwh", "Pot. Eigenverbrauch", "#287c71"),
    ]
    maximum = max(float(row.get(key, 0) or 0) for row in rows for key, _, _ in series)
    maximum = max(maximum, 1.0)

    draw.text((left, 30), "Monatliche Energieübersicht", fill="#17201d", font=title_font)
    for tick in range(6):
        value = maximum * tick / 5
        y = top + plot_height - plot_height * tick / 5
        draw.line((left, y, width - right, y), fill="#dfe6e2", width=2)
        label = _format_number(value)
        label_box = draw.textbbox((0, 0), label, font=small_font)
        draw.text((left - 16 - (label_box[2] - label_box[0]), y - 11), label, fill="#60716a", font=small_font)

    draw.line((left, top, left, top + plot_height), fill="#60716a", width=3)
    draw.line((left, top + plot_height, width - right, top + plot_height), fill="#60716a", width=3)
    draw.text((18, top - 42), "kWh", fill="#60716a", font=label_font)

    group_width = plot_width / len(rows)
    bar_width = min(44, max(10, int(group_width * 0.22)))
    gap = max(4, int(bar_width * 0.15))
    bars_width = len(series) * bar_width + (len(series) - 1) * gap
    for row_index, row in enumerate(rows):
        group_left = left + row_index * group_width + (group_width - bars_width) / 2
        for series_index, (key, _, color) in enumerate(series):
            value = float(row.get(key, 0) or 0)
            bar_height = plot_height * value / maximum
            x0 = group_left + series_index * (bar_width + gap)
            y0 = top + plot_height - bar_height
            draw.rectangle((x0, y0, x0 + bar_width, top + plot_height), fill=color)
        label = _month_label(row.get("timestamp"))
        label_box = draw.textbbox((0, 0), label, font=small_font)
        draw.text(
            (left + row_index * group_width + (group_width - (label_box[2] - label_box[0])) / 2, top + plot_height + 18),
            label,
            fill="#60716a",
            font=small_font,
        )

    legend_x = left
    for _, label, color in series:
        draw.rectangle((legend_x, height - 48, legend_x + 24, height - 24), fill=color)
        draw.text((legend_x + 34, height - 51), label, fill="#17201d", font=legend_font)
        legend_x += 34 + draw.textbbox((0, 0), label, font=legend_font)[2] + 36
    return _png_bytes(image)


def _residual_solar_chart(
    rows: list[dict[str, Any]],
    *,
    title: str,
    label_key: str,
    label_formatter: Any = str,
) -> BytesIO | None:
    if not rows:
        return None

    width, height = 1500, 760
    left, right, top, bottom = 120, 40, 120, 125
    plot_width = width - left - right
    plot_height = height - top - bottom
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(34, bold=True)
    label_font = _font(22)
    small_font = _font(18)
    maximum = max(float(row.get("residual_solar_kwh", 0) or 0) for row in rows)
    maximum = max(maximum, 1.0)

    draw.text((left, 30), title, fill="#17201d", font=title_font)
    for tick in range(6):
        value = maximum * tick / 5
        y = top + plot_height - plot_height * tick / 5
        draw.line((left, y, width - right, y), fill="#dfe6e2", width=2)
        label = _format_number(value)
        label_box = draw.textbbox((0, 0), label, font=small_font)
        draw.text((left - 16 - (label_box[2] - label_box[0]), y - 11), label, fill="#60716a", font=small_font)

    draw.line((left, top, left, top + plot_height), fill="#60716a", width=3)
    draw.line((left, top + plot_height, width - right, top + plot_height), fill="#60716a", width=3)
    draw.text((18, top - 42), "kWh", fill="#60716a", font=label_font)

    group_width = plot_width / len(rows)
    bar_width = min(54, max(10, int(group_width * 0.55)))
    for index, row in enumerate(rows):
        value = float(row.get("residual_solar_kwh", 0) or 0)
        bar_height = plot_height * value / maximum
        x0 = left + index * group_width + (group_width - bar_width) / 2
        y0 = top + plot_height - bar_height
        draw.rectangle((x0, y0, x0 + bar_width, top + plot_height), fill="#d79c24")
        label = label_formatter(row.get(label_key))
        label_box = draw.textbbox((0, 0), label, font=small_font)
        draw.text(
            (left + index * group_width + (group_width - (label_box[2] - label_box[0])) / 2, top + plot_height + 18),
            label,
            fill="#60716a",
            font=small_font,
        )

    draw.rectangle((left, height - 48, left + 24, height - 24), fill="#d79c24")
    draw.text((left + 34, height - 51), "Verbleibender Solarüberschuss", fill="#17201d", font=label_font)
    return _png_bytes(image)


def _residual_solar_heatmap(heatmap: dict[str, Any]) -> BytesIO | None:
    days = list(heatmap.get("days", []))
    hours = list(heatmap.get("hours", []))
    values = list(heatmap.get("values", []))
    if not days or not hours:
        return None

    width, height = 1800, 920
    left, right, top, bottom = 135, 65, 125, 125
    plot_width = width - left - right
    plot_height = height - top - bottom
    cell_width = plot_width / len(days)
    cell_height = plot_height / len(hours)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(34, bold=True)
    label_font = _font(22)
    small_font = _font(18)
    draw.text((left, 30), "Verbleibender Solarüberschuss nach Datum und Tageszeit", fill="#17201d", font=title_font)

    value_by_cell = {
        (str(item["date"]), int(item["hour"])): item.get("residual_solar_kwh")
        for item in values
        if int(item.get("interval_count", 0)) > 0
    }
    maximum = max(float(heatmap.get("scale_max_residual_solar_kwh", 0) or 0), 1.0)
    selected_peak = heatmap.get("selected_peak")
    low = (247, 244, 239)
    high = (155, 77, 85)

    def color(value: object) -> tuple[int, int, int]:
        if value is None:
            return (232, 237, 234)
        ratio = min(max(float(value) / maximum, 0.0), 1.0)
        return tuple(round(start + (end - start) * ratio) for start, end in zip(low, high))

    for day_index, day in enumerate(days):
        x0 = left + day_index * cell_width
        if str(day).endswith("-01"):
            draw.line((x0, top, x0, top + plot_height), fill="#b6c1bd", width=2)
            draw.text((x0 + 3, top + plot_height + 14), datetime.fromisoformat(str(day)).strftime("%b"), fill="#60716a", font=small_font)
        for hour_index, hour_label in enumerate(hours):
            hour = int(str(hour_label).split(":")[0])
            y0 = top + hour_index * cell_height
            draw.rectangle(
                (x0, y0, x0 + max(cell_width, 1), y0 + cell_height),
                fill=color(value_by_cell.get((str(day), hour))),
            )

    for hour_index, hour_label in enumerate(hours):
        if hour_index % 2 == 0 or hour_index == len(hours) - 1:
            y = top + hour_index * cell_height + cell_height / 2
            draw.text((left - 72, y - 11), str(hour_label), fill="#60716a", font=small_font)

    if selected_peak:
        peak_day_index = days.index(str(selected_peak["date"]))
        peak_hour_label = f"{int(selected_peak['hour']):02d}:00"
        peak_hour_index = hours.index(peak_hour_label)
        peak_x = left + (peak_day_index + 0.5) * cell_width
        peak_y = top + (peak_hour_index + 0.5) * cell_height
        draw.ellipse((peak_x - 10, peak_y - 10, peak_x + 10, peak_y + 10), outline="#17201d", width=4)
        peak_timestamp = f"{selected_peak['date']}T{peak_hour_label}"
        callout = (
            f"Maximum Auswahl: {_format_heatmap_kwh(selected_peak['residual_solar_kwh'])}\n"
            f"{_timestamp_label(peak_timestamp)}"
        )
        callout_box = draw.multiline_textbbox((0, 0), callout, font=small_font, spacing=4)
        callout_width = callout_box[2] - callout_box[0]
        callout_height = callout_box[3] - callout_box[1]
        callout_x = width - right - callout_width - 28
        callout_y = top + 14
        draw.line((peak_x, peak_y, callout_x - 8, callout_y + callout_height / 2), fill="#17201d", width=3)
        draw.rounded_rectangle(
            (callout_x - 10, callout_y - 8, callout_x + callout_width + 10, callout_y + callout_height + 8),
            radius=7,
            fill="white",
            outline="#9fb2aa",
            width=2,
        )
        draw.multiline_text((callout_x, callout_y), callout, fill="#17201d", font=small_font, spacing=4)

    legend_left = left
    legend_top = height - 55
    legend_width = 340
    for step in range(100):
        ratio = step / 99
        fill = tuple(round(start + (end - start) * ratio) for start, end in zip(low, high))
        x0 = legend_left + legend_width * ratio
        draw.rectangle((x0, legend_top, x0 + legend_width / 99 + 1, legend_top + 22), fill=fill)
    draw.text((legend_left, legend_top - 29), "Solarüberschuss pro Stunde", fill="#17201d", font=label_font)
    draw.text((legend_left, legend_top + 27), "0 kWh", fill="#60716a", font=small_font)
    maximum_label = _format_heatmap_kwh(maximum)
    maximum_box = draw.textbbox((0, 0), maximum_label, font=small_font)
    draw.text(
        (legend_left + legend_width - (maximum_box[2] - maximum_box[0]), legend_top + 27),
        maximum_label,
        fill="#60716a",
        font=small_font,
    )
    return _png_bytes(image)


def _arrow(draw: ImageDraw.ImageDraw, start: tuple[float, float], end: tuple[float, float], *, color: str) -> None:
    draw.line((*start, *end), fill=color, width=5)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    left = (end[0] - length * math.cos(angle - math.pi / 6), end[1] - length * math.sin(angle - math.pi / 6))
    right = (end[0] - length * math.cos(angle + math.pi / 6), end[1] - length * math.sin(angle + math.pi / 6))
    draw.polygon([end, left, right], fill=color)


def _flow_diagram(graph: dict[str, Any] | None) -> BytesIO | None:
    if not graph or not graph.get("nodes"):
        return None

    nodes = graph["nodes"]
    positions = {
        str(node["id"]): (
            float(node.get("position", {}).get("x", 0)),
            float(node.get("position", {}).get("y", 0)),
        )
        for node in nodes
    }
    min_x = min(position[0] for position in positions.values())
    min_y = min(position[1] for position in positions.values())
    max_x = max(position[0] for position in positions.values())
    max_y = max(position[1] for position in positions.values())
    scale = 1.35
    margin = 65
    node_width, node_height = 330, 148
    width = max(1100, int((max_x - min_x) * scale + node_width + margin * 2))
    height = max(500, int((max_y - min_y) * scale + node_height + margin * 2))
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    label_font = _font(22, bold=True)
    value_font = _font(21, bold=True)
    small_font = _font(18)
    edge_font = _font(17, bold=True)

    boxes: dict[str, tuple[float, float, float, float]] = {}
    for node_id, (x, y) in positions.items():
        x0 = margin + (x - min_x) * scale
        y0 = margin + (y - min_y) * scale
        boxes[node_id] = (x0, y0, x0 + node_width, y0 + node_height)

    for edge in graph.get("edges", []):
        source = boxes.get(str(edge.get("source")))
        target = boxes.get(str(edge.get("target")))
        if not source or not target:
            continue
        start = (source[2], (source[1] + source[3]) / 2)
        end = (target[0], (target[1] + target[3]) / 2)
        _arrow(draw, start, end, color="#7b8a84")
        label = str(edge.get("label") or "")
        if label:
            center = ((start[0] + end[0]) / 2, (start[1] + end[1]) / 2)
            label_box = draw.textbbox((0, 0), label, font=edge_font)
            padding = 6
            draw.rounded_rectangle(
                (
                    center[0] - (label_box[2] - label_box[0]) / 2 - padding,
                    center[1] - 15 - padding,
                    center[0] + (label_box[2] - label_box[0]) / 2 + padding,
                    center[1] + 15 + padding,
                ),
                radius=8,
                fill="white",
                outline="#dfe6e2",
            )
            draw.text(
                (center[0] - (label_box[2] - label_box[0]) / 2, center[1] - 14),
                label,
                fill="#60716a",
                font=edge_font,
            )

    fills = {
        "source": "#fff8e8",
        "zev-source": "#fff3cf",
        "zev-consumption": "#edf6ff",
        "zev-grid-import": "#f7f4ef",
        "zev-grid-export": "#f7f4ef",
        "solar": "#f2fbf5",
        "leg": "#eff7fb",
        "participant": "#ffffff",
        "grid": "#f7f4ef",
    }
    for node in nodes:
        node_id = str(node["id"])
        data = node.get("data", {})
        box = boxes[node_id]
        kind = str(data.get("kind") or "")
        draw.rounded_rectangle(box, radius=16, fill=fills.get(kind, "#ffffff"), outline="#9fb2aa", width=3)
        y = box[1] + 16
        for line in textwrap.wrap(str(data.get("label") or node_id), width=28)[:2]:
            draw.text((box[0] + 16, y), line, fill="#17201d", font=label_font)
            y += 27
        value = str(data.get("value") or "")
        if value:
            draw.text((box[0] + 16, y + 5), value, fill="#1f6f64", font=value_font)
            y += 31
        detail = str(data.get("subtitle") or data.get("tariff") or "")
        if detail:
            draw.text((box[0] + 16, y + 4), detail[:36], fill="#60716a", font=small_font)
    return _png_bytes(image)


def _add_picture(document: DocumentType, image: BytesIO, *, width_inches: float = 6.5) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(image, width=Inches(width_inches))


def create_project_word_report(
    *,
    project_id: str,
    db_path: str | Path,
    plant_cache_path: str | Path = DEFAULT_PLANT_CACHE,
    template_path: str | Path | None = None,
    network_discount: Decimal = Decimal("0.40"),
    leg_energy_price_rp_per_kwh: Decimal | None = None,
) -> WordReport:
    summary = summarize_project(project_id=project_id, db_path=db_path)
    calculation = calculate_leg_potential(
        project_id=project_id,
        network_discount=network_discount,
        leg_energy_price_rp_per_kwh=leg_energy_price_rp_per_kwh,
        db_path=db_path,
    )
    residual_solar = analyze_residual_solar(project_id=project_id, db_path=db_path)
    optimization = optimize_leg_combinations(
        project_id=project_id,
        network_discount=network_discount,
        leg_energy_price_rp_per_kwh=leg_energy_price_rp_per_kwh,
        db_path=db_path,
    )
    plants = match_project_plants(
        project_id=project_id,
        municipality=str(summary["municipality"]),
        db_path=db_path,
        cache_path=plant_cache_path,
    )
    totals = calculation.get("totals", {})
    rates = calculation.get("rates", {})
    document = _render_template(
        template_path=template_path,
        context={
            "project_name": str(summary["name"]),
            "municipality": str(summary["municipality"]),
            "report_date": date.today().strftime("%d.%m.%Y"),
            "consumption_kwh": _format_kwh(totals.get("consumption_kwh")),
            "available_solar_kwh": _format_kwh(totals.get("available_solar_kwh")),
            "local_solar_kwh": _format_kwh(totals.get("local_solar_consumed_kwh")),
            "remaining_grid_import_kwh": _format_kwh(totals.get("remaining_grid_import_kwh")),
            "local_solar_energy_value_chf": _format_chf(totals.get("local_solar_energy_value_chf")),
            "producer_revenue_chf": _format_chf(totals.get("producer_revenue_chf")),
            "energy_price_savings_chf": _format_chf(totals.get("energy_price_savings_chf")),
            "network_discount_chf": _format_chf(totals.get("network_discount_chf")),
            "total_monetary_benefit_chf": _format_chf(totals.get("total_monetary_benefit_chf")),
        },
    )

    document.add_heading("Methode", level=1)
    document.add_paragraph(
        "Die Potenzialanalyse basiert auf den hochgeladenen 15-Minuten-Messwerten. "
        "PV-Überschuss und Strombedarf werden für jedes Zeitintervall gegenübergestellt."
    )
    document.add_paragraph(
        "Der potenzielle lokale Eigenverbrauch entspricht pro Zeitintervall dem kleineren Wert aus verfügbarem "
        "PV-Überschuss und Bedarf. Die Monats- und Gesamtwerte werden aus diesen Intervallergebnissen aggregiert."
    )
    document.add_paragraph(
        f"Der vereinbarte lokale Energiepreis beträgt {float(rates.get('leg_energy_price_rp_per_kwh', 0)):.2f} Rp./kWh; "
        f"der externe Energiepreis {float(rates.get('external_energy_rp_per_kwh', 0)):.2f} Rp./kWh. "
        f"Zusätzlich werden {float(rates.get('network_discount', 0)):.0%} des CKW-Netznutzungseinheitstarifs "
        f"von {float(rates.get('discountable_netznutzung_rp_per_kwh', 0)):.2f} Rp./kWh rabattiert. "
        "Systemdienstleistungen, Stromreserve, solidarisierte Kosten, Leistungspreise, Messung, öffentliche "
        "Abgaben und MWST sind nicht Bestandteil des LEG-Rabatts."
    )
    document.add_paragraph(
        "Die Optimierung vergleicht die möglichen Kombinationen der hochgeladenen Verbrauchsteilnehmer. "
        "Bei grossen Gemeinschaften wird die vollständige Teilnehmergruppe als angenommene optimale LEG bewertet."
    )
    participant_prices = calculation.get("participants", [])
    if participant_prices:
        document.add_heading("LEG-Preise je Teilnehmer", level=2)
        _add_table(
            document,
            headers=[
                "Teilnehmer",
                "Verbrauch",
                "Lokaler Solarstrom",
                "Energie extern",
                "Energie lokal",
                "Netztarif",
                "Rabatt je kWh",
                "Verbrauchervorteil",
            ],
            rows=[
                [
                    str(item["label"]),
                    _format_kwh(item["measured_consumption_kwh"]),
                    _format_kwh(item["allocated_local_solar_kwh"]),
                    f"{item['external_energy_rp_per_kwh']:.2f} Rp./kWh",
                    f"{item['leg_energy_price_rp_per_kwh']:.2f} Rp./kWh",
                    f"{item['discountable_netznutzung_rp_per_kwh']:.2f} Rp./kWh",
                    f"{item['network_discount_rp_per_kwh']:.2f} Rp./kWh",
                    _format_chf(item["consumer_total_savings_chf"]),
                ]
                for item in participant_prices
            ],
        )

    document.add_heading("Produktionseinheiten", level=1)
    document.add_paragraph(
        f"Erfasste PV-Leistung: {_format_number(plants.get('total_installed_kwp'), 1)} kWp. "
        f"Geschätzte maximale Anschlussleistung: {_format_number(plants.get('max_connection_ampere'))} A."
    )
    plant_matches = plants.get("matches", [])
    if plant_matches:
        _add_table(
            document,
            headers=["Adresse", "Gemeinde", "Installierte Leistung", "Anschlussleistung", "Quelle"],
            rows=[
                [
                    str(item["address"]),
                    str(item["municipality"]),
                    f"{_format_number(item['installed_kwp'], 1)} kWp",
                    f"{_format_number(item['max_connection_ampere'])} A",
                    str(item["confidence"]),
                ]
                for item in plant_matches
            ],
        )
    else:
        document.add_paragraph("Keine Produktionseinheiten wurden dem Projekt zugeordnet.")

    document.add_heading("VZEV-Potenzial", level=1)

    monthly_solar = {
        str(item["timestamp"]): item
        for item in project_timeseries(project_id=project_id, kind="solar", resolution="month", db_path=db_path)
    }
    monthly_rows = [
        {
            **item,
            "pv_excess_kwh": monthly_solar.get(str(item["timestamp"]), {}).get("pv_excess", 0),
        }
        for item in calculation.get("monthly", [])
    ]
    if monthly_rows:
        document.add_heading("Monatliche Energieübersicht", level=1)
        document.add_paragraph(
            "Die Monatswerte zeigen den verfügbaren PV-Überschuss, den Bedarf und den daraus potenziell "
            "lokal nutzbaren Solarstrom."
        )
        _add_table(
            document,
            headers=["Monat", "PV-Überschuss gemessen", "Bedarf", "Pot. Eigenverbrauch"],
            rows=[
                [
                    _month_label(item.get("timestamp")),
                    _format_kwh(item.get("pv_excess_kwh")),
                    _format_kwh(item.get("consumption_kwh")),
                    _format_kwh(item.get("local_solar_kwh")),
                ]
                for item in monthly_rows
            ],
        )
        chart = _monthly_energy_chart(monthly_rows)
        if chart:
            _add_picture(document, chart)

    residual_totals = residual_solar["totals"]
    document.add_heading("Verbleibender Solarüberschuss", level=1)
    document.add_paragraph(
        "Verbleibender Solarüberschuss entsteht in einem 15-Minuten-Intervall, wenn der verfügbare Solarstrom "
        "grösser ist als der gleichzeitig durch die LEG nutzbare Bedarf."
    )
    document.add_paragraph(
        f"Insgesamt verbleiben {_format_kwh(residual_totals['residual_solar_kwh'])} beziehungsweise "
        f"{_format_pct(residual_totals['residual_solar_pct'])} des verfügbaren Solarstroms. "
        f"Überschuss tritt in {_format_number(residual_totals['residual_interval_count'])} von "
        f"{_format_number(residual_totals['interval_count'])} Intervallen auf "
        f"({_format_pct(residual_totals['residual_interval_pct'])})."
    )
    peak_interval = residual_solar.get("peak_interval")
    if peak_interval:
        document.add_paragraph(
            f"Der höchste verbleibende Überschuss innerhalb eines 15-Minuten-Intervalls beträgt "
            f"{_format_kwh(peak_interval['residual_solar_kwh'])} am {_timestamp_label(peak_interval['timestamp'])}. "
            "In diesem Intervall stehen "
            f"{_format_kwh(peak_interval['available_solar_kwh'])} Solarstrom einem gleichzeitigen Bedarf von "
            f"{_format_kwh(peak_interval['consumption_kwh'])} gegenüber."
        )
    else:
        document.add_paragraph("In den ausgewerteten Messwerten verbleibt kein Solarüberschuss.")

    residual_monthly = residual_solar.get("monthly", [])
    if residual_monthly:
        document.add_heading("Verteilung nach Monat", level=2)
        _add_table(
            document,
            headers=["Monat", "Solarüberschuss", "Anteil Solarstrom", "Intervalle mit Überschuss", "Anteil Intervalle"],
            rows=[
                [
                    _month_label(item["timestamp"]),
                    _format_kwh(item["residual_solar_kwh"]),
                    _format_pct(item["residual_solar_pct"]),
                    _format_number(item["residual_interval_count"]),
                    _format_pct(item["residual_interval_pct"]),
                ]
                for item in residual_monthly
            ],
        )
        chart = _residual_solar_chart(
            residual_monthly,
            title="Verbleibender Solarüberschuss nach Monat",
            label_key="timestamp",
            label_formatter=_month_label,
        )
        if chart:
            _add_picture(document, chart)

    residual_hourly = residual_solar.get("hourly", [])
    if residual_hourly:
        document.add_heading("Verteilung nach Datum und Tageszeit", level=2)
        _add_table(
            document,
            headers=["Uhrzeit", "Solarüberschuss total", "Ø je Überschussintervall", "Maximum", "Intervalle mit Überschuss"],
            rows=[
                [
                    str(item["label"]),
                    _format_kwh(item["residual_solar_kwh"]),
                    _format_kwh(item["average_residual_when_present_kwh"]),
                    _format_kwh(item["peak_residual_solar_kwh"]),
                    _format_number(item["residual_interval_count"]),
                ]
                for item in residual_hourly
            ],
        )
        chart = _residual_solar_heatmap(residual_solar.get("heatmap", {}))
        if chart:
            _add_picture(document, chart)

    peak_intervals = residual_solar.get("peak_intervals", [])
    if peak_intervals:
        document.add_heading("Höchste 15-Minuten-Intervalle", level=2)
        _add_table(
            document,
            headers=["Zeitpunkt", "Verfügbarer Solarstrom", "Gleichzeitiger Bedarf", "Verbleibender Überschuss"],
            rows=[
                [
                    _timestamp_label(item["timestamp"]),
                    _format_kwh(item["available_solar_kwh"]),
                    _format_kwh(item["consumption_kwh"]),
                    _format_kwh(item["residual_solar_kwh"]),
                ]
                for item in peak_intervals
            ],
        )

    document.add_heading("Importierte Dateien", level=1)
    _add_table(
        document,
        headers=["Datei", "Upload-Typ", "Messfluss", "Energiemenge"],
        rows=[
            [
                str(item["original_filename"]),
                str(item["upload_kind"]),
                str(item["flow_kind"]),
                _format_kwh(item["total_kwh"]),
            ]
            for item in summary["files"]
        ],
    )

    best = optimization.get("best")
    if best:
        document.add_heading("Optimale LEG-Zusammensetzung", level=1)
        document.add_paragraph(f"Optimierter Verbrauchervorteil: {_format_chf(best.get('cost_savings_chf'))}")
        document.add_paragraph(f"Produzentenerlös für lokale Energie: {_format_chf(best.get('producer_revenue_chf'))}")
        flow_diagram = _flow_diagram(best.get("flow"))
        if flow_diagram:
            document.add_paragraph(
                "Die Grafik visualisiert dieselben optimierten Energieflüsse wie die interaktive Darstellung im Portal."
            )
            _add_picture(document, flow_diagram)
        _add_table(
            document,
            headers=[
                "Teilnehmer",
                "Tarif",
                "Jahresverbrauch",
                "Solarstrom",
                "Netzrabatt",
                "Energiepreisvorteil",
                "Produzentenerlös",
                "Verbrauchervorteil",
            ],
            rows=[
                [
                    str(item["label"]),
                    str(item["tariff_tier"]),
                    _format_kwh(item["annualized_kwh"]),
                    _format_kwh(item["allocated_local_solar_kwh"]),
                    _format_chf(item["leg_network_discount_chf"]),
                    _format_chf(item["energy_price_savings_chf"]),
                    _format_chf(item["producer_revenue_chf"]),
                    _format_chf(item["cost_savings_chf"]),
                ]
                for item in best["members"]
            ],
        )
    elif summary.get("zev_flow"):
        document.add_heading("Gemessene CKW-ZEV-Energieflüsse", level=1)
        flow_diagram = _flow_diagram(summary["zev_flow"])
        if flow_diagram:
            _add_picture(document, flow_diagram)

    warnings = sorted(
        set(
            summary.get("warnings", [])
            + calculation.get("warnings", [])
            + optimization.get("warnings", [])
            + plants.get("warnings", [])
        )
    )
    if warnings:
        document.add_heading("Hinweise", level=1)
        for warning in warnings:
            document.add_paragraph(str(warning), style="List Bullet")

    assumptions = calculation.get("assumptions", []) + optimization.get("assumptions", [])
    if assumptions:
        document.add_heading("Annahmen", level=1)
        for assumption in assumptions:
            document.add_paragraph(str(assumption), style="List Bullet")

    output = BytesIO()
    document.save(output)
    return WordReport(
        filename=f"leg-potenzial-{_slug(summary['name'])}.docx",
        content=output.getvalue(),
    )
