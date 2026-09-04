from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any
import unicodedata

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

from .portal_analysis import (
    DEFAULT_PORTAL_DB,
    FLOW_CONSUMPTION,
    FLOW_GRID_IMPORT,
    FLOW_LOCAL_PV_CONSUMPTION,
    FLOW_PV_EXCESS,
    FLOW_PV_PRODUCTION,
    FLOW_ZEV_AGGREGATE,
    _deduped_file_cte,
    _leg_input_timeseries_for_project,
    calculate_leg_potential,
    connect_portal_db,
    get_project,
    optimize_leg_combinations,
    summarize_project,
)


EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FLOW_LABELS = {
    FLOW_CONSUMPTION: "Verbrauch",
    FLOW_PV_PRODUCTION: "PV-Produktion",
    FLOW_PV_EXCESS: "PV-Ueberschuss",
    FLOW_GRID_IMPORT: "Netzbezug",
    FLOW_LOCAL_PV_CONSUMPTION: "Eigenverbrauch",
    FLOW_ZEV_AGGREGATE: "ZEV-Aggregat",
}


@dataclass(frozen=True)
class ExcelExport:
    filename: str
    content: bytes


def _slug(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "leg-validierung"))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text or "leg-validierung"


def _decode_json(value: object, default: Any) -> Any:
    try:
        return json.loads(str(value or ""))
    except json.JSONDecodeError:
        return default


def _cell_value(value: object) -> object:
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value)
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime().replace(tzinfo=None)
    if isinstance(value, np.generic):
        return value.item()
    if value is pd.NaT:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def _ordered_year_index(timestamps: pd.Series | pd.Index) -> pd.DatetimeIndex:
    parsed = pd.to_datetime(pd.Series(timestamps), errors="coerce").dropna()
    if parsed.empty:
        return pd.DatetimeIndex([], name="timestamp")
    years = sorted(int(year) for year in parsed.dt.year.unique())
    ranges = [
        pd.date_range(
            start=pd.Timestamp(year=year, month=1, day=1),
            end=pd.Timestamp(year=year + 1, month=1, day=1) - pd.Timedelta(minutes=15),
            freq="15min",
        )
        for year in years
    ]
    return pd.DatetimeIndex(np.concatenate([item.to_numpy() for item in ranges]), name="timestamp")


def _unique_labels(values: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    labels: list[str] = []
    for value in values:
        base = value[:220]
        count = seen.get(base, 0) + 1
        seen[base] = count
        labels.append(base if count == 1 else f"{base} ({count})")
    return labels


def _metadata_label(row: pd.Series) -> str:
    metadata = _decode_json(row.get("metadata_json"), {})
    return str(metadata.get("vertrag") or row.get("original_filename") or row.get("file_id"))


def _file_rows(files: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for _, row in files.iterrows():
        metadata = _decode_json(row.get("metadata_json"), {})
        warnings = _decode_json(row.get("warnings_json"), [])
        rows.append(
            {
                "file_id": row.get("file_id"),
                "original_filename": row.get("original_filename"),
                "upload_kind": row.get("upload_kind"),
                "flow_kind": row.get("flow_kind"),
                "label": metadata.get("vertrag") or row.get("original_filename"),
                "address": metadata.get("adresse"),
                "zaehlpunkt": metadata.get("zaehlpunkt"),
                "energy_column": row.get("energy_column"),
                "total_kwh": row.get("total_kwh"),
                "warnings": warnings,
                "metadata_json": metadata,
            }
        )
    return pd.DataFrame(rows)


def _collect_frames(
    *,
    project_id: str,
    db_path: str | Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    con = connect_portal_db(db_path)
    try:
        files = con.execute(
            f"""
            WITH selected_files AS ({_deduped_file_cte()})
            SELECT
                f.file_id,
                f.original_filename,
                f.upload_kind,
                f.flow_kind,
                f.energy_column,
                f.total_kwh,
                f.metadata_json,
                f.warnings_json,
                f.ingested_at
            FROM files f
            JOIN selected_files sf USING (file_id)
            WHERE f.project_id = ?
            ORDER BY f.ingested_at, f.original_filename, f.file_id
            """,
            [project_id, project_id],
        ).df()
        measurements = con.execute(
            f"""
            WITH selected_files AS ({_deduped_file_cte()})
            SELECT
                m.timestamp,
                m.file_id,
                f.original_filename,
                f.flow_kind AS file_flow_kind,
                m.flow_kind AS measurement_flow_kind,
                ANY_VALUE(m.meter_label) AS meter_label,
                ANY_VALUE(m.address) AS address,
                ANY_VALUE(m.zaehlpunkt) AS zaehlpunkt,
                SUM(m.kwh) AS kwh
            FROM measurements m
            JOIN selected_files sf USING (file_id)
            JOIN files f USING (file_id)
            WHERE m.project_id = ?
            GROUP BY 1, 2, 3, 4, 5
            ORDER BY 1, 2, 5
            """,
            [project_id, project_id],
        ).df()
        leg_inputs = _leg_input_timeseries_for_project(con, project_id=project_id)
    finally:
        con.close()
    return files, measurements, leg_inputs


def _raw_timeseries_frame(
    *,
    files: pd.DataFrame,
    measurements: pd.DataFrame,
    full_index: pd.DatetimeIndex,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    if measurements.empty:
        return pd.DataFrame({"timestamp": full_index}), pd.DataFrame(
            columns=[
                "series_column",
                "file_id",
                "original_filename",
                "file_flow_kind",
                "measurement_flow_kind",
                "meter_label",
                "address",
                "zaehlpunkt",
            ]
        )

    labels: list[str] = []
    file_label_by_id = {
        str(row["file_id"]): _metadata_label(row)
        for _, row in files.iterrows()
    }
    for _, row in measurements[["file_id", "measurement_flow_kind"]].drop_duplicates().iterrows():
        file_id = str(row["file_id"])
        flow_kind = str(row["measurement_flow_kind"])
        labels.append(f"{file_label_by_id.get(file_id, file_id)} | {FLOW_LABELS.get(flow_kind, flow_kind)} kWh")
    keys = [
        f"{row['file_id']}|{row['measurement_flow_kind']}"
        for _, row in measurements[["file_id", "measurement_flow_kind"]].drop_duplicates().iterrows()
    ]
    label_by_key = dict(zip(keys, _unique_labels(labels), strict=True))

    source = measurements.copy()
    source["timestamp"] = pd.to_datetime(source["timestamp"], errors="raise")
    source["series_key"] = source["file_id"].astype(str) + "|" + source["measurement_flow_kind"].astype(str)
    wide = source.pivot_table(index="timestamp", columns="series_key", values="kwh", aggfunc="sum").sort_index()
    wide = wide.reindex(full_index)
    wide = wide.rename(columns=label_by_key).reset_index()

    mapping_rows: list[dict[str, object]] = []
    first_rows = (
        source.sort_values(["series_key", "timestamp"])
        .drop_duplicates("series_key")
        .set_index("series_key")
    )
    for key, label in label_by_key.items():
        row = first_rows.loc[key]
        mapping_rows.append(
            {
                "series_column": label,
                "file_id": row["file_id"],
                "original_filename": row["original_filename"],
                "file_flow_kind": row["file_flow_kind"],
                "measurement_flow_kind": row["measurement_flow_kind"],
                "meter_label": row["meter_label"],
                "address": row["address"],
                "zaehlpunkt": row["zaehlpunkt"],
            }
        )
    return wide, pd.DataFrame(mapping_rows)


def _leg_interval_frame(
    *,
    leg_inputs: pd.DataFrame,
    full_index: pd.DatetimeIndex,
    calculation: dict[str, Any],
) -> pd.DataFrame:
    columns = [
        "timestamp",
        "consumption_kwh",
        "available_solar_kwh",
        "local_solar_kwh",
        "residual_solar_kwh",
        "grid_import_kwh",
        "external_energy_rp_per_kwh",
        "leg_energy_price_rp_per_kwh",
        "discountable_netznutzung_rp_per_kwh",
        "network_discount",
        "producer_revenue_chf",
        "energy_price_savings_chf",
        "network_discount_chf",
        "consumer_total_savings_chf",
    ]
    if full_index.empty:
        return pd.DataFrame(columns=columns)

    rates = calculation.get("rates") or {}
    source = leg_inputs.copy()
    if source.empty:
        source = pd.DataFrame(index=full_index, columns=["consumption_kwh", "available_solar_kwh"])
    else:
        source["timestamp"] = pd.to_datetime(source["timestamp"], errors="raise")
        source = source.set_index("timestamp").reindex(full_index)
    source = source.fillna(0.0)
    frame = pd.DataFrame(
        {
            "timestamp": full_index,
            "consumption_kwh": pd.to_numeric(source["consumption_kwh"], errors="coerce").fillna(0.0).to_numpy(),
            "available_solar_kwh": pd.to_numeric(source["available_solar_kwh"], errors="coerce").fillna(0.0).to_numpy(),
            "local_solar_kwh": None,
            "residual_solar_kwh": None,
            "grid_import_kwh": None,
            "external_energy_rp_per_kwh": rates.get("external_energy_rp_per_kwh"),
            "leg_energy_price_rp_per_kwh": rates.get("leg_energy_price_rp_per_kwh"),
            "discountable_netznutzung_rp_per_kwh": rates.get("discountable_netznutzung_rp_per_kwh"),
            "network_discount": rates.get("network_discount"),
            "producer_revenue_chf": None,
            "energy_price_savings_chf": None,
            "network_discount_chf": None,
            "consumer_total_savings_chf": None,
        }
    )
    return frame


def _style_header(ws, row: int) -> None:
    fill = PatternFill("solid", fgColor="DDE9E4")
    for cell in ws[row]:
        cell.font = Font(bold=True, color="17201D")
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _append_dataframe(ws, df: pd.DataFrame, *, start_row: int = 1) -> int:
    row_index = start_row
    for col_index, column in enumerate(df.columns, start=1):
        ws.cell(row=row_index, column=col_index, value=str(column))
    _style_header(ws, row_index)
    for values in df.itertuples(index=False, name=None):
        row_index += 1
        for col_index, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col_index, value=_cell_value(value))
    return row_index + 2


def _append_section_title(ws, title: str, row: int) -> int:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=13, color="17201D")
    return row + 1


def _append_key_values(ws, title: str, values: dict[str, object], *, start_row: int) -> int:
    row = _append_section_title(ws, title, start_row)
    ws.cell(row=row, column=1, value="key")
    ws.cell(row=row, column=2, value="value")
    _style_header(ws, row)
    row += 1
    for key, value in values.items():
        ws.cell(row=row, column=1, value=str(key))
        ws.cell(row=row, column=2, value=_cell_value(value))
        row += 1
    return row + 2


def _write_leg_interval_sheet(ws, df: pd.DataFrame) -> None:
    _append_dataframe(ws, df, start_row=1)
    max_row = ws.max_row
    if max_row < 2:
        return
    for row in range(2, max_row + 1):
        ws.cell(row=row, column=4, value=f"=MIN(B{row},C{row})")
        ws.cell(row=row, column=5, value=f"=MAX(C{row}-D{row},0)")
        ws.cell(row=row, column=6, value=f"=MAX(B{row}-D{row},0)")
        ws.cell(row=row, column=11, value=f"=D{row}*H{row}/100")
        ws.cell(row=row, column=12, value=f"=D{row}*(G{row}-H{row})/100")
        ws.cell(row=row, column=13, value=f"=D{row}*I{row}*J{row}/100")
        ws.cell(row=row, column=14, value=f"=L{row}+M{row}")


def _set_sheet_defaults(ws, *, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    for column_index in range(1, ws.max_column + 1):
        letter = get_column_letter(column_index)
        header = ws.cell(row=1, column=column_index).value
        if str(header).lower() == "timestamp":
            ws.column_dimensions[letter].width = 20
            for row_index in range(2, ws.max_row + 1):
                ws.cell(row=row_index, column=column_index).number_format = "yyyy-mm-dd hh:mm"
        else:
            ws.column_dimensions[letter].width = min(max(len(str(header or "")) + 2, 12), 42)


def _summary_sheet(
    ws,
    *,
    project: dict[str, Any],
    summary: dict[str, Any],
    calculation: dict[str, Any],
    export_started_at: datetime,
) -> None:
    row = _append_key_values(
        ws,
        "Projekt",
        {
            "project_id": project.get("project_id"),
            "name": project.get("name"),
            "municipality": project.get("municipality"),
            "created_at": project.get("created_at"),
            "exported_at": export_started_at,
            "purpose": "Validierung der Portalberechnungen in Excel",
        },
        start_row=1,
    )
    row = _append_key_values(ws, "Summen aus Messdaten", summary.get("totals", {}), start_row=row)
    row = _append_key_values(ws, "Tarife und Faktoren", calculation.get("rates", {}), start_row=row)
    _append_key_values(ws, "Berechnete LEG-Summen", calculation.get("totals", {}), start_row=row)


def _monthly_sheet(ws, *, calculation: dict[str, Any]) -> None:
    row = _append_key_values(ws, "Tarife und Faktoren", calculation.get("rates", {}), start_row=1)
    row = _append_key_values(ws, "Berechnete LEG-Summen", calculation.get("totals", {}), start_row=row)
    monthly = pd.DataFrame(calculation.get("monthly") or [])
    if not monthly.empty:
        row = _append_section_title(ws, "Monatliche LEG-Rechnung", row)
        _append_dataframe(ws, monthly, start_row=row)


def _optimization_sheet(ws, *, optimization: dict[str, Any]) -> None:
    row = _append_key_values(
        ws,
        "Optimierung",
        {
            "tariff_year": optimization.get("tariff_year"),
            "network_discount": optimization.get("network_discount"),
            "annualization_factor": optimization.get("annualization_factor"),
            "participant_count": optimization.get("participant_count"),
            "evaluated_participant_count": optimization.get("evaluated_participant_count"),
            "combination_count": optimization.get("combination_count"),
            "evaluated_all": optimization.get("evaluated_all"),
            "optimization_mode": optimization.get("optimization_mode"),
        },
        start_row=1,
    )
    best = optimization.get("best") or {}
    if best:
        scalar_best = {
            key: value
            for key, value in best.items()
            if key not in {"members", "solar_sources", "flow", "participant_ids", "participant_labels"}
        }
        row = _append_key_values(ws, "Beste Kombination", scalar_best, start_row=row)
        row = _append_section_title(ws, "Teilnehmer der besten Kombination", row)
        row = _append_dataframe(ws, pd.DataFrame(best.get("members") or []), start_row=row)
        row = _append_section_title(ws, "Solarquellen der besten Kombination", row)
        row = _append_dataframe(ws, pd.DataFrame(best.get("solar_sources") or []), start_row=row)

    combinations = pd.DataFrame(optimization.get("combinations") or [])
    if not combinations.empty:
        row = _append_section_title(ws, "Top-Kombinationen", row)
        _append_dataframe(ws, combinations, start_row=row)


def _assumptions_sheet(
    ws,
    *,
    summary: dict[str, Any],
    calculation: dict[str, Any],
    optimization: dict[str, Any],
) -> None:
    rows: list[dict[str, object]] = []
    for source, values in (
        ("summary_warning", summary.get("warnings") or []),
        ("calculation_warning", calculation.get("warnings") or []),
        ("optimization_warning", optimization.get("warnings") or []),
        ("calculation_assumption", calculation.get("assumptions") or []),
        ("optimization_assumption", optimization.get("assumptions") or []),
    ):
        rows.extend({"source": source, "text": value} for value in values)
    rows.append(
        {
            "source": "excel_export_note",
            "text": (
                "Messreihen 15min zeigt leere Zellen, wenn kein importierter Messwert fuer die Serie vorliegt. "
                "Die LEG Rechnung 15min behandelt fehlende Werte als 0 kWh und nutzt Excel-Formeln pro Intervall."
            ),
        }
    )
    _append_dataframe(ws, pd.DataFrame(rows), start_row=1)


def create_project_excel_export(
    *,
    project_id: str,
    tariff_year: int | None = None,
    network_discount: Decimal = Decimal("0.40"),
    leg_energy_price_rp_per_kwh: Decimal | None = None,
    db_path: str | Path = DEFAULT_PORTAL_DB,
) -> ExcelExport:
    export_started_at = datetime.now()
    project = get_project(project_id=project_id, db_path=db_path)
    summary = summarize_project(project_id=project_id, db_path=db_path)
    calculation = calculate_leg_potential(
        project_id=project_id,
        tariff_year=tariff_year,
        network_discount=network_discount,
        leg_energy_price_rp_per_kwh=leg_energy_price_rp_per_kwh,
        db_path=db_path,
    )
    optimization = optimize_leg_combinations(
        project_id=project_id,
        tariff_year=tariff_year,
        network_discount=network_discount,
        leg_energy_price_rp_per_kwh=leg_energy_price_rp_per_kwh,
        db_path=db_path,
    )
    files, measurements, leg_inputs = _collect_frames(project_id=project_id, db_path=db_path)

    timestamp_sources: list[pd.Series | pd.Index] = []
    if not measurements.empty:
        timestamp_sources.append(measurements["timestamp"])
    if not leg_inputs.empty:
        timestamp_sources.append(leg_inputs["timestamp"])
    full_index = (
        _ordered_year_index(pd.concat([pd.Series(source) for source in timestamp_sources], ignore_index=True))
        if timestamp_sources
        else pd.DatetimeIndex([], name="timestamp")
    )
    raw_timeseries, series_mapping = _raw_timeseries_frame(
        files=files,
        measurements=measurements,
        full_index=full_index,
    )
    leg_interval = _leg_interval_frame(
        leg_inputs=leg_inputs,
        full_index=full_index,
        calculation=calculation,
    )

    workbook = Workbook()
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    ws_project = workbook.active
    ws_project.title = "Projekt"
    _summary_sheet(
        ws_project,
        project=project,
        summary=summary,
        calculation=calculation,
        export_started_at=export_started_at,
    )

    ws_files = workbook.create_sheet("Dateien")
    row = _append_section_title(ws_files, "Importierte Dateien", 1)
    row = _append_dataframe(ws_files, _file_rows(files), start_row=row)
    row = _append_section_title(ws_files, "Messreihen-Spalten", row)
    _append_dataframe(ws_files, series_mapping, start_row=row)

    ws_raw = workbook.create_sheet("Messreihen 15min")
    _append_dataframe(ws_raw, raw_timeseries, start_row=1)

    ws_leg = workbook.create_sheet("LEG Rechnung 15min")
    _write_leg_interval_sheet(ws_leg, leg_interval)

    ws_monthly = workbook.create_sheet("Monate Summen")
    _monthly_sheet(ws_monthly, calculation=calculation)

    ws_optimization = workbook.create_sheet("Optimierung")
    _optimization_sheet(ws_optimization, optimization=optimization)

    ws_assumptions = workbook.create_sheet("Annahmen Warnungen")
    _assumptions_sheet(
        ws_assumptions,
        summary=summary,
        calculation=calculation,
        optimization=optimization,
    )

    for worksheet in workbook.worksheets:
        _set_sheet_defaults(worksheet)

    output = BytesIO()
    workbook.save(output)
    return ExcelExport(
        filename=f"leg-validierung-{_slug(project.get('name'))}.xlsx",
        content=output.getvalue(),
    )
