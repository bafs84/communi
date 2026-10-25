from __future__ import annotations

"""Editable, presentation-oriented overview for one monthly vZEV bill."""

from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .vzev_invoice import allocate_ckw_invoice_totals
from .billing_tariff import invoice_tariff_label


GREEN = "1F6F64"
DARK = "17312C"
LIGHT_GREEN = "E8F2EF"
LIGHT_GREY = "F2F5F2"
AMBER = "FFF1D6"
WHITE = "FFFFFF"
MONEY_QUANT = Decimal("0.01")
ENERGY_QUANT = Decimal("0.001")


def _money(value: float | Decimal) -> str:
    return f"CHF {float(value):,.2f}".replace(",", "'")


def _kwh(value: float | Decimal) -> str:
    return f"{float(value):,.3f} kWh".replace(",", "'")


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(8.5)


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _shade(table.rows[0].cells[index], GREEN)
        _set_cell(table.rows[0].cells[index], header, bold=True, color=WHITE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            if row_index % 2:
                _shade(cells[column_index], LIGHT_GREY)
            _set_cell(cells[column_index], value)


def _add_kpis(document: Document, values: list[tuple[str, str, str]]) -> None:
    table = document.add_table(rows=1, cols=len(values))
    for index, (label, value, fill) in enumerate(values):
        cell = table.cell(0, index)
        _shade(cell, fill)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = paragraph.add_run(label.upper() + "\n")
        label_run.bold = True
        label_run.font.size = Pt(7.5)
        label_run.font.color.rgb = RGBColor.from_string(DARK)
        value_run = paragraph.add_run(value)
        value_run.bold = True
        value_run.font.size = Pt(15)
        value_run.font.color.rgb = RGBColor.from_string(GREEN)


def _producer_allocations(parsed, result: dict[str, Any]) -> list[dict[str, object]]:
    feed = result["feed_in_compensation"]
    production_rows: list[dict[str, object]] = []
    for producer in parsed.producers:
        production_rows.append(
            {
                "producer": producer.display_name,
                "metering_code": producer.metering_code,
                "production_kwh": Decimal(str(float(parsed.dataframe[producer.column_name].sum()))),
            }
        )
    total_production = sum((row["production_kwh"] for row in production_rows), Decimal("0"))
    local_energy_total = Decimal(str(result["community"]["community_local_pv_kwh"]))
    measured_export_energy_total = Decimal(str(result["community"]["community_excess_total_kwh"]))
    export_energy_total = Decimal(str(feed["billing_excess_pv_kwh"]))
    local_total = Decimal(str(result["community"]["community_local_energy_chf"]))
    export_total = Decimal(str(feed["total_chf"]))
    allocated_local_energy = Decimal("0")
    allocated_measured_export_energy = Decimal("0")
    allocated_export_energy = Decimal("0")
    allocated_local = Decimal("0")
    allocated_export = Decimal("0")
    for row in production_rows:
        production = row["production_kwh"]
        share = production / total_production if total_production else Decimal("0")
        local_energy = (local_energy_total * share).quantize(ENERGY_QUANT, rounding=ROUND_HALF_UP)
        measured_export_energy = (measured_export_energy_total * share).quantize(
            ENERGY_QUANT, rounding=ROUND_HALF_UP
        )
        export_energy = (export_energy_total * share).quantize(ENERGY_QUANT, rounding=ROUND_HALF_UP)
        local = (local_total * share).quantize(MONEY_QUANT, rounding=ROUND_HALF_UP)
        export = (export_total * share).quantize(MONEY_QUANT, rounding=ROUND_HALF_UP)
        row["share"] = share
        row["local_energy_kwh"] = local_energy
        row["measured_export_energy_kwh"] = measured_export_energy
        row["export_energy_kwh"] = export_energy
        row["local_payout_chf"] = local
        row["export_payout_chf"] = export
        row["total_payout_chf"] = local + export
        allocated_local_energy += local_energy
        allocated_measured_export_energy += measured_export_energy
        allocated_export_energy += export_energy
        allocated_local += local
        allocated_export += export
    if production_rows:
        last = production_rows[-1]
        local_energy_adjustment = local_energy_total - allocated_local_energy
        measured_export_energy_adjustment = measured_export_energy_total - allocated_measured_export_energy
        export_energy_adjustment = export_energy_total - allocated_export_energy
        local_adjustment = local_total - allocated_local
        export_adjustment = export_total - allocated_export
        last["local_energy_kwh"] = last["local_energy_kwh"] + local_energy_adjustment
        last["measured_export_energy_kwh"] = (
            last["measured_export_energy_kwh"] + measured_export_energy_adjustment
        )
        last["export_energy_kwh"] = last["export_energy_kwh"] + export_energy_adjustment
        last["local_payout_chf"] = last["local_payout_chf"] + local_adjustment
        last["export_payout_chf"] = last["export_payout_chf"] + export_adjustment
        last["total_payout_chf"] = last["local_payout_chf"] + last["export_payout_chf"]
    return production_rows


def create_community_overview_docx(
    *,
    result: dict[str, Any],
    parsed,
    issuer: str,
) -> bytes:
    feed = result["feed_in_compensation"]
    aggregate = result["aggregated_invoice"]
    community = result["community"]
    rates = result["rates"]
    allocations = allocate_ckw_invoice_totals(result)
    producer_rows = _producer_allocations(parsed, result)

    participant_total = sum(
        Decimal(str(row["participant_invoice_total_chf"])) for row in allocations.values()
    )
    local_pv_total = Decimal(str(community["community_local_energy_chf"]))
    export_credit = Decimal(str(feed["total_chf"]))
    net_ckw = Decimal(str(feed["net_ckw_payable_chf"]))
    producer_payout = local_pv_total + export_credit
    control_balance = participant_total - net_ckw - producer_payout

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    document.core_properties.title = "vZEV Gemeinschaftsübersicht"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("vZEV Gemeinschaftsübersicht")
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor.from_string(GREEN)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"{issuer} · {result['period_start']} bis {result['period_end']} · Tarifjahr {result['tariff_year']}"
    )
    document.add_paragraph(f"CKW-Vertragstarif: {invoice_tariff_label(result)}")

    _add_kpis(
        document,
        [
            ("CKW Bruttobelastung", _money(aggregate["invoice_total_chf"]), LIGHT_GREEN),
            ("Rückliefergutschrift", f"− {_money(export_credit)}", AMBER),
            ("Zahlbetrag an CKW", _money(net_ckw), LIGHT_GREEN),
            ("Teilnehmerrechnungen", _money(participant_total), LIGHT_GREY),
        ],
    )

    document.add_heading("So fliesst das Geld", level=1)
    _add_table(
        document,
        ["Position", "Betrag", "Bedeutung"],
        [
            ["Teilnehmerrechnungen", _money(participant_total), "CKW-Kostenanteile plus lokal genutzte PV-Energie"],
            ["Zahlbetrag an CKW", _money(net_ckw), "Bruttobelastung abzüglich Rückliefergutschrift und inkl. Schlussrundung"],
            ["Auszahlung an PV-Betreiber", _money(producer_payout), "Lokale PV-Energie plus CKW-Gutschrift für Export und HKN"],
            ["Kontrollsaldo", _money(control_balance), "Soll nach vollständiger Verteilung CHF 0.00 betragen"],
        ],
    )

    document.add_heading("CKW-Kosten und Rücklieferung", level=1)
    _add_table(
        document,
        ["Kostenblock", "Menge", "Tarif", "Betrag"],
        [
            ["Energielieferung", _kwh(aggregate["billing_grid_import_kwh"]), f"{rates['energy_rp_per_kwh']:.2f} Rp./kWh", _money(aggregate["energy_chf"])],
            ["Netznutzung Einheitstarif", _kwh(aggregate["billing_grid_import_kwh"]), f"{rates['grid_unit_rp_per_kwh']:.2f} Rp./kWh", _money(aggregate["grid_unit_chf"])],
            ["Systemdienstleistungen Swissgrid", _kwh(aggregate["billing_grid_import_kwh"]), f"{rates['grid_swissgrid_rp_per_kwh']:.2f} Rp./kWh", _money(aggregate["grid_swissgrid_chf"])],
            ["Stromreserve", _kwh(aggregate["billing_grid_import_kwh"]), f"{rates['grid_stromreserve_rp_per_kwh']:.2f} Rp./kWh", _money(aggregate["grid_stromreserve_chf"])],
            ["Solidarisierte Kosten", _kwh(aggregate["billing_grid_import_kwh"]), f"{rates['grid_solidarity_rp_per_kwh']:.2f} Rp./kWh", _money(aggregate["grid_solidarisierte_kosten_chf"])],
            ["Leistungspreis", f"{aggregate['billing_peak_kw']:.1f} kW", f"{rates['peak_price_chf_per_kw_month']:.2f} CHF/kW-Monat", _money(aggregate["peak_price_chf"])],
            ["Blindenergie", f"{aggregate['billing_blindenergie_kvarh']:.0f} kvarh", f"{rates['blindenergie_rp_per_kvarh']:.2f} Rp./kvarh", _money(aggregate["blindenergie_chf"])],
            ["Direkte Messungen", str(aggregate["metering_direct_count"]), f"{rates['direct_meter_chf_per_month']:.2f} CHF/Monat", _money(aggregate["metering_direct_chf"])],
            ["Indirekte Messungen", str(aggregate["metering_indirect_count"]), f"{rates['indirect_meter_chf_per_month']:.2f} CHF/Monat", _money(aggregate["metering_indirect_chf"])],
            ["Virtuelle Messungen", str(aggregate["metering_virtual_count"]), f"{rates['virtual_meter_chf_per_month']:.2f} CHF/Monat", _money(aggregate["metering_virtual_chf"])],
            ["Netzzuschlag", _kwh(aggregate["billing_grid_import_kwh"]), f"{rates['netzzuschlag_rp_per_kwh']:.2f} Rp./kWh", _money(aggregate["public_netzzuschlag_chf"])],
            ["Gemeindeabgabe", _kwh(aggregate["billing_grid_import_kwh"]), f"{rates['gemeinde_abgabe_rp_per_kwh']:.2f} Rp./kWh", _money(aggregate["public_gemeinde_chf"])],
            ["MWST Belastung", "", f"{rates['vat_rate']:.1%}" if rates.get("vat_rate") is not None else "", _money(aggregate["vat_chf"])],
            ["CKW Bruttobelastung", "", "", _money(aggregate["invoice_total_chf"])],
            ["Energie-Rückliefervergütung", _kwh(feed["billing_excess_pv_kwh"]), f"{feed['energy_rp_per_kwh']:.3f} Rp./kWh", f"− {_money(feed['energy_revenue_chf'])}"],
            ["HKN-Vergütung", _kwh(feed["billing_excess_pv_kwh"]), f"{feed['hkn_rp_per_kwh']:.2f} Rp./kWh", f"− {_money(feed['hkn_revenue_chf'])}"],
            ["MWST Rückliefergutschrift", "", f"{rates['vat_rate']:.1%}" if rates.get("vat_rate") is not None else "", f"− {_money(feed['vat_chf'])}"],
            ["Schlussrundung", "", "", _money(feed["payment_rounding_chf"])],
            ["Zahlbetrag an CKW", "", "", _money(net_ckw)],
        ],
    )

    document.add_heading("Aufteilung auf die Teilnehmer", level=1)
    participant_energy_rows: list[list[str]] = []
    participant_cost_rows: list[list[str]] = []
    participant_price_rows: list[list[str]] = []
    participant_consumption_total = Decimal("0")
    participant_grid_total = Decimal("0")
    participant_local_total = Decimal("0")
    for consumer in result["consumers"]:
        key = str(consumer["metering_code"] or consumer["display_name"])
        allocation = allocations[key]
        participant_consumption_total += Decimal(str(consumer["consumption_kwh"]))
        participant_grid_total += Decimal(str(consumer["allocated_grid_import_kwh"]))
        participant_local_total += Decimal(str(consumer["allocated_local_pv_kwh"]))
        participant_energy_rows.append(
            [
                str(consumer["display_name"]),
                _kwh(consumer["consumption_kwh"]),
                _kwh(consumer["allocated_grid_import_kwh"]),
                _kwh(consumer["allocated_local_pv_kwh"]),
            ]
        )
        participant_cost_rows.append(
            [
                str(consumer["display_name"]),
                _money(allocation["ckw_invoice_share_chf"]),
                _money(allocation["local_energy_chf"]),
                _money(allocation["participant_invoice_total_chf"]),
            ]
        )
        participant_price_rows.append(
            [
                str(consumer["display_name"]),
                f"{consumer['standard_product_reference_rp_per_kwh']:.2f} Rp./kWh",
                f"{consumer['local_energy_factor']:.0%}",
                f"{consumer['local_energy_rp_per_kwh']:.2f} Rp./kWh",
                _kwh(consumer["allocated_local_pv_kwh"]),
                _money(consumer["local_energy_chf"]),
            ]
        )
    participant_energy_rows.append(
        [
            "Summe",
            _kwh(participant_consumption_total),
            _kwh(participant_grid_total),
            _kwh(participant_local_total),
        ]
    )
    participant_cost_rows.append(
        ["Summe", _money(net_ckw + export_credit), _money(local_pv_total), _money(participant_total)]
    )
    document.add_heading("Energiemengen", level=2)
    _add_table(
        document,
        ["Teilnehmer", "Verbrauch", "CKW-Netzbezug", "Lokale PV"],
        participant_energy_rows,
    )
    document.add_heading("Preis der lokalen PV-Energie", level=2)
    _add_table(
        document,
        ["Teilnehmer", "Standardprodukt inkl. MWST", "Faktor", "PV-Tarif inkl. MWST", "Lokale PV", "Betrag"],
        participant_price_rows,
    )
    document.add_heading("Kostenaufteilung", level=2)
    document.add_paragraph(
        "CKW-Kostenanteil: Bezugskosten inkl. MWST und Rundung. "
        "Rückliefergutschriften werden separat bei der Auszahlung an die PV-Betreiber berücksichtigt."
    )
    _add_table(
        document,
        ["Teilnehmer", "CKW-Kostenanteil", "Lokale PV", "Rechnung total"],
        participant_cost_rows,
    )

    document.add_heading("Auszahlung an die PV-Betreiber", level=1)
    document.add_heading("PV-Energiemengen", level=2)
    _add_table(
        document,
        ["PV-Anlage", "Produktion", "Lokal genutzt", "Export gemessen", "Export abgerechnet"],
        [
            [
                str(row["producer"]),
                _kwh(row["production_kwh"]),
                _kwh(row["local_energy_kwh"]),
                _kwh(row["measured_export_energy_kwh"]),
                _kwh(row["export_energy_kwh"]),
            ]
            for row in producer_rows
        ]
        + [
            [
                "Summe",
                _kwh(sum((row["production_kwh"] for row in producer_rows), Decimal("0"))),
                _kwh(community["community_local_pv_kwh"]),
                _kwh(community["community_excess_total_kwh"]),
                _kwh(feed["billing_excess_pv_kwh"]),
            ]
        ],
    )
    document.add_heading("PV-Auszahlungen", level=2)
    local_average_rate = (
        float(local_pv_total) * 100 / float(community["community_local_pv_kwh"])
        if community["community_local_pv_kwh"]
        else 0.0
    )
    export_average_rate = (
        float(export_credit) * 100 / float(feed["billing_excess_pv_kwh"])
        if feed["billing_excess_pv_kwh"]
        else 0.0
    )
    _add_table(
        document,
        ["PV-Anlage", "Lokale PV", "Ø Preis", "CHF lokal", "Export", "Preis inkl. HKN/MWST", "CHF Export", "Total"],
        [
            [
                str(row["producer"]),
                _kwh(row["local_energy_kwh"]),
                f"{local_average_rate:.2f} Rp./kWh",
                _money(row["local_payout_chf"]),
                _kwh(row["export_energy_kwh"]),
                f"{export_average_rate:.3f} Rp./kWh",
                _money(row["export_payout_chf"]),
                _money(row["total_payout_chf"]),
            ]
            for row in producer_rows
        ]
        + [[
            "Summe",
            _kwh(community["community_local_pv_kwh"]),
            f"{local_average_rate:.2f} Rp./kWh",
            _money(local_pv_total),
            _kwh(feed["billing_excess_pv_kwh"]),
            f"{export_average_rate:.3f} Rp./kWh",
            _money(export_credit),
            _money(producer_payout),
        ]],
    )

    document.add_heading("Berechnungsgrundlagen", level=1)
    assumptions = [
        (
            f"Interner PV-Preis: je Teilnehmer {result['rates']['local_energy_factor']:.0%} des vollständigen "
            "CKW-Standardproduktpreises inklusive Energie, Netznutzung, Leistung, Messung, Abgaben und MWST. "
            f"Der mengengewichtete PV-Preis beträgt {result['rates']['local_energy_rp_per_kwh']:.2f} Rp./kWh."
        ),
        f"Netzbezug: {_kwh(aggregate['billing_grid_import_kwh'])}; Monatsspitze: {aggregate['billing_peak_kw']:.1f} kW.",
        (
            f"Rücklieferung: {_kwh(community['community_excess_total_kwh'])} gemessen, "
            f"{_kwh(feed['billing_excess_pv_kwh'])} von CKW abgerechnet zu "
            f"{feed['energy_rp_per_kwh']:.3f} Rp./kWh; "
            f"HKN {feed['hkn_rp_per_kwh']:.2f} Rp./kWh."
        ),
        f"Quelle Energievergütung: {feed['energy_rate_source']}.",
        f"Quelle HKN: {feed['hkn_rate_source']}.",
        (
            "Die Rückliefergutschrift wird mangels viertelstündlicher Erzeugerzuordnung proportional zur "
            "gemessenen PV-Produktion aufgeteilt."
        ),
        (
            "Die Teilnehmerrechnungen enthalten keinen hypothetischen Vergleich ohne vZEV. "
            "Sie zeigen ausschliesslich die aktuelle Abrechnungsperiode."
        ),
    ]
    if feed.get("warning"):
        assumptions.append(f"Hinweis: {feed['warning']}")
    for assumption in assumptions:
        document.add_paragraph(assumption, style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
