from __future__ import annotations

"""Editable participant-by-month overview for a vZEV invoice batch."""

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from .billing_tariff import TARIFF_LABELS, require_invoice_tariff_tier


GREEN = "1F6F64"
LIGHT_GREY = "F2F5F2"
WHITE = "FFFFFF"
MONTH_NAMES = (
    "",
    "Januar",
    "Februar",
    "März",
    "April",
    "Mai",
    "Juni",
    "Juli",
    "August",
    "September",
    "Oktober",
    "November",
    "Dezember",
)


def _money(value: float) -> str:
    return f"CHF {value:,.2f}".replace(",", "'")


def _kwh(value: float) -> str:
    return f"{value:,.3f} kWh".replace(",", "'")


def _period_label(period_key: str) -> str:
    year = int(period_key[:4])
    month = int(period_key[4:6])
    return f"{MONTH_NAMES[month]} {year}"


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_participant_table(
    document: Document,
    *,
    participant: dict[str, Any],
    period_keys: tuple[str, ...],
) -> None:
    monthly_values = {row["period_key"]: row for row in participant["months"]}
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    headers = ("Monat", "Verbrauch", "Geld an PV", "Geld an CKW", "Total")
    widths = (Cm(4.2), Cm(4.5), Cm(4.5), Cm(4.5), Cm(4.5))
    for index, (header, width) in enumerate(zip(headers, widths, strict=True)):
        table.columns[index].width = width
        cell = table.rows[0].cells[index]
        cell.width = width
        _shade(cell, GREEN)
        _set_cell(cell, header, bold=True, color=WHITE)

    for row_index, period_key in enumerate(period_keys):
        values = monthly_values.get(period_key)
        row = table.add_row()
        row_values = (
            (
                _period_label(period_key),
                _kwh(float(values["consumption_kwh"])),
                _money(float(values["pv_payment_chf"])),
                _money(float(values["ckw_payment_chf"])),
                _money(float(values["total_chf"])),
            )
            if values is not None
            else (_period_label(period_key), "–", "–", "–", "–")
        )
        for column_index, (value, width) in enumerate(zip(row_values, widths, strict=True)):
            cell = row.cells[column_index]
            cell.width = width
            if row_index % 2:
                _shade(cell, LIGHT_GREY)
            _set_cell(cell, value)

    total_row = table.add_row()
    total_values = (
        "Summe",
        _kwh(float(participant["totals"]["consumption_kwh"])),
        _money(float(participant["totals"]["pv_payment_chf"])),
        _money(float(participant["totals"]["ckw_payment_chf"])),
        _money(float(participant["totals"]["total_chf"])),
    )
    for column_index, (value, width) in enumerate(zip(total_values, widths, strict=True)):
        cell = total_row.cells[column_index]
        cell.width = width
        _shade(cell, LIGHT_GREY)
        _set_cell(cell, value, bold=True)


def create_multi_month_overview_docx(
    *,
    period_keys: tuple[str, ...],
    participants: list[dict[str, Any]],
    issuer: str,
    tariff_tier: str,
    tariff_years: tuple[int, ...],
) -> bytes:
    """Create one readable table per participant across every batch month."""

    tariff_tier = require_invoice_tariff_tier(tariff_tier)
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    document.core_properties.title = "vZEV Mehrmonatsübersicht"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("vZEV Mehrmonatsübersicht")
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor.from_string(GREEN)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"{issuer} · {_period_label(period_keys[0])} bis {_period_label(period_keys[-1])}"
    )
    document.add_paragraph(
        f"CKW-Vertragstarif: {TARIFF_LABELS[tariff_tier]} (vertraglich gewählt) · "
        f"Tarifjahre: {', '.join(str(year) for year in tariff_years)}"
    )
    explanation = document.add_paragraph()
    explanation.add_run("Geld an PV").bold = True
    explanation.add_run(" zeigt die interne Vergütung für lokal genutzte PV-Energie. ")
    explanation.add_run("Geld an CKW").bold = True
    explanation.add_run(
        " zeigt die Bezugskosten inklusive MWST und Rechnungsabgleich. "
        "Rückliefergutschriften werden separat bei der Auszahlung an die PV-Betreiber berücksichtigt. "
        "Beide Beträge ergeben das Total des Teilnehmers."
    )

    for participant in participants:
        heading = document.add_heading(str(participant["participant"]), level=1)
        heading.paragraph_format.keep_with_next = True
        metering_code = str(participant.get("metering_code") or "")
        if metering_code:
            paragraph = document.add_paragraph(f"Messpunkt: {metering_code}")
            paragraph.paragraph_format.keep_with_next = True
        _add_participant_table(
            document,
            participant=participant,
            period_keys=period_keys,
        )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
