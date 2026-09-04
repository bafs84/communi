from __future__ import annotations

"""Editable Word report for the complete vZEV planning workspace."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

from .portal_analysis import (
    DEFAULT_PLANT_CACHE,
    analyze_battery_placements,
    analyze_load_shifting,
    match_project_plants,
    project_timeseries,
    simulate_project_vzev,
    summarize_project,
)
from .portal_word_report import (
    WordReport,
    _add_picture,
    _add_table,
    _flow_diagram,
    _font,
    _format_kwh,
    _format_number,
    _format_pct,
    _png_bytes,
    _slug,
)
from .vzev_simulation import DEFAULT_VZEV_SIMULATION_TARIFF_YEAR, VzevSimulation


def _format_chf(value: object) -> str:
    return f"CHF {_format_number(value, 2)}"


def _format_rp(value: object) -> str:
    return f"{_format_number(value, 2)} Rp./kWh"


def _format_kw(value: object) -> str:
    return f"{_format_number(value, 1)} kW"


def _format_cost_change(value: object) -> str:
    number = float(value or 0)
    sign = "+" if number > 0.004 else "−" if number < -0.004 else ""
    return f"{sign}{_format_chf(abs(number))}"


def _tariff_tier_label(value: object) -> str:
    if value == "gt_50000":
        return "> 50 MWh/Jahr"
    if value == "lt_50000":
        return "≤ 50 MWh/Jahr"
    return str(value or "–")


def _month_label(value: object) -> str:
    text = str(value or "")
    try:
        return datetime.strptime(text[:7], "%Y-%m").strftime("%m.%Y")
    except ValueError:
        return text


def _configure_document(document: DocumentType) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = RGBColor(31, 111, 100)
    for section in document.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)


def _shade_cell(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _style_tables(document: DocumentType) -> None:
    for table in document.tables:
        if not table.rows:
            continue
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        header_properties.append(repeat_header)
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index == 0:
                    _shade_cell(cell, "1F6F64")
                elif row_index % 2 == 0:
                    _shade_cell(cell, "F1F6F4")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Aptos"
                        run.font.size = Pt(8.5)
                        if row_index == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)


def _add_key_values(document: DocumentType, rows: list[tuple[str, str]]) -> None:
    _add_table(document, headers=["Eingabe", "Wert"], rows=[[label, value] for label, value in rows])


def _add_bullets(document: DocumentType, values: list[object]) -> None:
    for value in values:
        document.add_paragraph(str(value), style="List Bullet")


def _series_chart(
    rows: list[dict[str, Any]],
    *,
    title: str,
    series: list[tuple[str, str, str]],
    label_key: str = "timestamp",
) -> BytesIO | None:
    if not rows:
        return None

    width, height = 1500, 760
    left, right, top, bottom = 125, 45, 125, 125
    plot_width = width - left - right
    plot_height = height - top - bottom
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(34, bold=True)
    label_font = _font(21)
    small_font = _font(18)
    maximum = max(float(row.get(key, 0) or 0) for row in rows for key, _, _ in series)
    maximum = max(maximum, 1.0)

    draw.text((left, 30), title, fill="#17201d", font=title_font)
    for tick in range(6):
        value = maximum * tick / 5
        y = top + plot_height - plot_height * tick / 5
        draw.line((left, y, width - right, y), fill="#dfe6e2", width=2)
        label = _format_number(value)
        box = draw.textbbox((0, 0), label, font=small_font)
        draw.text((left - 14 - box[2] + box[0], y - 10), label, fill="#60716a", font=small_font)
    draw.line((left, top, left, top + plot_height), fill="#60716a", width=3)
    draw.line((left, top + plot_height, width - right, top + plot_height), fill="#60716a", width=3)
    draw.text((22, top - 42), "kWh", fill="#60716a", font=label_font)

    group_width = plot_width / len(rows)
    bar_width = min(36, max(3, int(group_width * 0.72 / len(series))))
    for row_index, row in enumerate(rows):
        group_x = left + row_index * group_width + group_width / 2
        for series_index, (key, _, color) in enumerate(series):
            value = float(row.get(key, 0) or 0)
            bar_height = plot_height * value / maximum
            offset = (series_index - (len(series) - 1) / 2) * bar_width
            x0 = group_x + offset - bar_width / 2
            draw.rectangle((x0, top + plot_height - bar_height, x0 + bar_width, top + plot_height), fill=color)
        if len(rows) <= 18 or row_index % max(1, len(rows) // 12) == 0:
            label = _month_label(row.get(label_key))
            box = draw.textbbox((0, 0), label, font=small_font)
            draw.text((group_x - (box[2] - box[0]) / 2, top + plot_height + 16), label, fill="#60716a", font=small_font)

    legend_x = left
    legend_y = height - 48
    for _, label, color in series:
        draw.rectangle((legend_x, legend_y, legend_x + 22, legend_y + 22), fill=color)
        draw.text((legend_x + 31, legend_y - 3), label, fill="#17201d", font=label_font)
        legend_x += draw.textbbox((0, 0), label, font=label_font)[2] + 75
    return _png_bytes(image)


def _load_shifting_heatmap(heatmap: dict[str, Any]) -> BytesIO | None:
    days = list(heatmap.get("days", []))
    hours = list(heatmap.get("hours", []))
    if not days or not hours:
        return None

    width, height = 1800, 920
    left, right, top, bottom = 135, 65, 125, 125
    plot_width = width - left - right
    plot_height = height - top - bottom
    cell_width = plot_width / len(days)
    cell_height = plot_height / len(hours)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(34, bold=True)
    label_font = _font(22)
    small_font = _font(18)
    draw.text((left, 30), "Stündliches Verbrauchsmuster über das Jahr", fill="#17201d", font=title_font)

    value_by_cell = {
        (str(item["date"]), int(item["hour"])): item.get("demand_kwh")
        for item in heatmap.get("values", [])
        if int(item.get("interval_count", 0)) > 0
    }
    maximum = max(float(heatmap.get("scale_max_demand_kwh", 0) or 0), 1.0)
    low = (239, 246, 244)
    high = (27, 111, 100)

    def color(value: object) -> tuple[int, int, int]:
        if value is None:
            return (232, 237, 234)
        ratio = min(max(float(value) / maximum, 0.0), 1.0)
        return tuple(round(start + (end - start) * ratio) for start, end in zip(low, high))

    for day_index, day in enumerate(days):
        x0 = left + day_index * cell_width
        if str(day).endswith("-01"):
            draw.line((x0, top, x0, top + plot_height), fill="#b6c1bd", width=2)
            draw.text((x0 + 3, top + plot_height + 14), datetime.fromisoformat(str(day)).strftime("%b"), fill="#60716a", font=small_font)
        for hour_index, hour_label in enumerate(hours):
            hour = int(str(hour_label).split(":")[0])
            y0 = top + hour_index * cell_height
            draw.rectangle((x0, y0, x0 + max(cell_width, 1), y0 + cell_height), fill=color(value_by_cell.get((str(day), hour))))

    for hour_index, hour_label in enumerate(hours):
        if hour_index % 2 == 0 or hour_index == len(hours) - 1:
            y = top + hour_index * cell_height + cell_height / 2
            draw.text((left - 72, y - 11), str(hour_label), fill="#60716a", font=small_font)

    selected_peak = heatmap.get("selected_peak")
    if selected_peak:
        peak_day = str(selected_peak["date"])
        peak_hour_label = f"{int(selected_peak['hour']):02d}:00"
        if peak_day in days and peak_hour_label in hours:
            peak_x = left + (days.index(peak_day) + 0.5) * cell_width
            peak_y = top + (hours.index(peak_hour_label) + 0.5) * cell_height
            draw.ellipse((peak_x - 10, peak_y - 10, peak_x + 10, peak_y + 10), outline="#17201d", width=4)

    legend_left = left
    legend_top = height - 55
    legend_width = 340
    for step in range(100):
        ratio = step / 99
        fill = tuple(round(start + (end - start) * ratio) for start, end in zip(low, high))
        x0 = legend_left + legend_width * ratio
        draw.rectangle((x0, legend_top, x0 + legend_width / 99 + 1, legend_top + 22), fill=fill)
    draw.text((legend_left, legend_top - 29), "Verbrauch pro Stunde", fill="#17201d", font=label_font)
    draw.text((legend_left, legend_top + 27), "0 kWh", fill="#60716a", font=small_font)
    maximum_label = _format_kwh(maximum)
    box = draw.textbbox((0, 0), maximum_label, font=small_font)
    draw.text((legend_left + legend_width - box[2] + box[0], legend_top + 27), maximum_label, fill="#60716a", font=small_font)
    return _png_bytes(image)


def _battery_chart(rows: list[dict[str, Any]], *, mode: str) -> BytesIO | None:
    if mode == "pv":
        configured = [row for row in rows if row.get("configured")]
        return _horizontal_bar_chart(
            configured,
            title="Zusätzlich lokal nutzbarer Solarstrom mit Batterie",
            value_key="additional_local_solar_kwh",
            value_formatter=_format_kwh,
            color="#d79c24",
        )
    return _horizontal_bar_chart(
        rows,
        title="Leistungstarif-Einsparung durch Peak Shaving",
        value_key="demand_charge_savings_chf",
        value_formatter=_format_chf,
        color="#287c71",
    )


def _horizontal_bar_chart(
    rows: list[dict[str, Any]],
    *,
    title: str,
    value_key: str,
    value_formatter: Any,
    color: str,
) -> BytesIO | None:
    rows = rows[:12]
    if not rows:
        return None
    width = 1500
    height = max(520, 155 + len(rows) * 72)
    left, right, top, bottom = 410, 190, 105, 45
    plot_width = width - left - right
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(34, bold=True)
    label_font = _font(20)
    value_font = _font(19, bold=True)
    draw.text((45, 28), title, fill="#17201d", font=title_font)
    maximum = max(float(row.get(value_key, 0) or 0) for row in rows)
    maximum = max(maximum, 1.0)
    available_height = height - top - bottom
    band = available_height / len(rows)
    for index, row in enumerate(rows):
        value = float(row.get(value_key, 0) or 0)
        y = top + index * band + band * 0.18
        bar_height = band * 0.62
        label = str(row.get("label") or row.get("file_id") or "")
        if len(label) > 38:
            label = f"{label[:35]}..."
        label_box = draw.textbbox((0, 0), label, font=label_font)
        draw.text((left - 18 - label_box[2] + label_box[0], y + 5), label, fill="#17201d", font=label_font)
        draw.rectangle((left, y, left + plot_width, y + bar_height), fill="#eef2f0")
        bar_width = plot_width * value / maximum
        draw.rectangle((left, y, left + bar_width, y + bar_height), fill=color)
        draw.text((left + bar_width + 12, y + 5), value_formatter(value), fill="#17201d", font=value_font)
    return _png_bytes(image)


def _vzev_cost_change_chart(rows: list[dict[str, Any]]) -> BytesIO | None:
    rows = rows[:20]
    if not rows:
        return None
    width = 1500
    height = max(560, 170 + len(rows) * 62)
    left, right, top, bottom = 405, 170, 125, 55
    plot_width = width - left - right
    plot_height = height - top - bottom
    zero_x = left + plot_width / 2
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(34, bold=True)
    label_font = _font(19)
    value_font = _font(18, bold=True)
    draw.text((45, 30), "Gesamtkostenänderung je Messpunkt", fill="#17201d", font=title_font)
    draw.text((45, 78), "Minus = Kostenreduktion · Plus = Mehrkosten", fill="#60716a", font=label_font)

    maximum = max(abs(float(row.get("total_cost_change_chf", 0) or 0)) for row in rows)
    maximum = max(maximum, 1.0)
    band = plot_height / len(rows)
    draw.line((zero_x, top - 10, zero_x, top + plot_height + 5), fill="#60716a", width=3)
    for index, row in enumerate(rows):
        value = float(row.get("total_cost_change_chf", 0) or 0)
        y = top + index * band + band * 0.18
        bar_height = band * 0.62
        bar_width = (plot_width / 2 - 20) * abs(value) / maximum
        x0 = zero_x if value >= 0 else zero_x - bar_width
        x1 = zero_x + bar_width if value >= 0 else zero_x
        color = "#a53b3b" if value > 0 else "#287c71"
        draw.rectangle((x0, y, x1, y + bar_height), fill=color)
        label = str(row.get("participant") or row.get("metering_code") or "")
        if len(label) > 38:
            label = f"{label[:35]}..."
        label_box = draw.textbbox((0, 0), label, font=label_font)
        draw.text((left - 18 - label_box[2] + label_box[0], y + 4), label, fill="#17201d", font=label_font)
        value_label = _format_cost_change(value)
        if value >= 0:
            draw.text((x1 + 10, y + 4), value_label, fill=color, font=value_font)
        else:
            value_box = draw.textbbox((0, 0), value_label, font=value_font)
            draw.text((x0 - 10 - value_box[2] + value_box[0], y + 4), value_label, fill=color, font=value_font)
    return _png_bytes(image)


def _simulation_monthly_energy(simulation: VzevSimulation) -> list[dict[str, Any]]:
    frame = simulation.dataframe.copy()
    if frame.empty:
        return []
    frame["month"] = frame["Zeitstempel"].astype(str).str.slice(0, 7)
    grouped = frame.groupby("month", sort=True).agg(
        consumption_kwh=("Energiebedarf_total_kWh", "sum"),
        pv_production_kwh=("Produktion_kWh", "sum"),
        local_pv_kwh=("Produktion_total_kWh", "sum"),
        grid_import_kwh=("Bezug_total_kWh", "sum"),
        pv_excess_kwh=("Rueckspeissung_total_kWh", "sum"),
    )
    return [{"month": month, **row} for month, row in grouped.to_dict(orient="index").items()]


def create_vzev_word_report(
    *,
    project_id: str,
    db_path: str | Path,
    plant_cache_path: str | Path = DEFAULT_PLANT_CACHE,
    tariff_year: int = DEFAULT_VZEV_SIMULATION_TARIFF_YEAR,
    consumption_file_id: str | None = None,
    load_shifting_year: int | None = None,
    fixed_local_energy_rp_per_kwh: Decimal | None = None,
) -> WordReport:
    summary = summarize_project(project_id=project_id, db_path=db_path)
    plants = match_project_plants(
        project_id=project_id,
        municipality=str(summary["municipality"]),
        db_path=db_path,
        cache_path=plant_cache_path,
    )
    solar_monthly = project_timeseries(project_id=project_id, kind="solar", resolution="month", db_path=db_path)
    load_shifting = analyze_load_shifting(
        project_id=project_id,
        consumption_file_id=consumption_file_id,
        year=load_shifting_year,
        db_path=db_path,
    )
    battery = analyze_battery_placements(project_id=project_id, tariff_year=tariff_year, db_path=db_path)
    simulation: VzevSimulation | None = None
    simulation_error: str | None = None
    try:
        simulation = simulate_project_vzev(
            project_id=project_id,
            tariff_year=tariff_year,
            fixed_local_energy_rp_per_kwh=fixed_local_energy_rp_per_kwh,
            db_path=db_path,
        )
    except ValueError as exc:
        simulation_error = str(exc)

    document = Document()
    _configure_document(document)
    document.core_properties.title = f"vZEV-Simulationsbericht – {summary['name']}"
    document.add_heading("vZEV-Simulationsbericht", level=0)
    document.add_paragraph(f"Projekt: {summary['name']}")
    document.add_paragraph(f"Gemeinde: {summary['municipality']}")
    document.add_paragraph(f"Berichtsdatum: {date.today().strftime('%d.%m.%Y')}")
    document.add_paragraph(f"CKW-Tarifjahr der finanziellen Simulation: {tariff_year}")
    document.add_paragraph(
        "Der Bericht dokumentiert die Eingaben, Energieflüsse, Ergebnisse, Diagramme und Berechnungsannahmen "
        "des vZEV-Arbeitsbereichs. Tabellen und Texte bleiben in Word editierbar."
    )

    document.add_heading("Übersicht", level=1)
    totals = summary.get("totals", {})
    _add_key_values(
        document,
        [
            ("Verbrauch gemessen", _format_kwh(totals.get("consumption_kwh"))),
            ("PV-Produktion gemessen", _format_kwh(totals.get("pv_production_kwh"))),
            ("PV-Überschuss gemessen", _format_kwh(totals.get("pv_excess_kwh"))),
            ("Netzbezug aus CKW-ZEV-Aggregat", _format_kwh(totals.get("grid_import_kwh"))),
            ("Lokaler PV-Verbrauch aus CKW-ZEV-Aggregat", _format_kwh(totals.get("local_pv_consumption_kwh"))),
            ("Zugeordnete PV-Leistung", f"{_format_number(plants.get('total_installed_kwp'), 1)} kWp"),
            ("Geschätzte Anschlussleistung", f"{_format_number(plants.get('max_connection_ampere'))} A"),
        ],
    )

    document.add_heading("Eingabedaten", level=1)
    document.add_paragraph(
        "Die eigentlichen 15-Minuten-Rohwerte bleiben in den importierten Quelldateien. Der Bericht weist jede "
        "verwendete Datei, ihre erkannte Rolle und die aggregierte Energiemenge aus."
    )
    if summary.get("files"):
        _add_table(
            document,
            headers=["Datei", "Upload-Typ", "Erkannter Messfluss", "Energiemenge"],
            rows=[
                [
                    str(item["original_filename"]),
                    str(item["upload_kind"]),
                    str(item["flow_kind"]),
                    _format_kwh(item["total_kwh"]),
                ]
                for item in summary["files"]
            ],
        )
    else:
        document.add_paragraph("Noch keine Messdatei importiert.")

    document.add_heading("PV-Anlagen und Anschluss", level=2)
    if plants.get("matches"):
        _add_table(
            document,
            headers=["Adresse", "Gemeinde", "Leistung", "Anschluss", "Zuordnung"],
            rows=[
                [
                    str(item["address"]),
                    str(item["municipality"]),
                    f"{_format_number(item['installed_kwp'], 1)} kWp",
                    f"{_format_number(item['max_connection_ampere'])} A",
                    str(item["confidence"]),
                ]
                for item in plants["matches"]
            ],
        )
    else:
        document.add_paragraph("Keine PV-Anlage wurde dem Projekt zugeordnet.")

    document.add_heading("Solarenergie", level=1)
    if solar_monthly:
        _add_table(
            document,
            headers=["Monat", "PV-Produktion", "PV-Überschuss", "Verfügbarer Solarstrom"],
            rows=[
                [
                    _month_label(item.get("timestamp")),
                    _format_kwh(item.get("pv_production")),
                    _format_kwh(item.get("pv_excess")),
                    _format_kwh(item.get("available_solar_kwh")),
                ]
                for item in solar_monthly
            ],
        )
        chart = _series_chart(
            solar_monthly,
            title="Solarenergie nach Monat",
            series=[
                ("pv_production", "PV-Produktion", "#d79c24"),
                ("pv_excess", "PV-Überschuss", "#9b4d55"),
                ("available_solar_kwh", "Verfügbarer Solarstrom", "#287c71"),
            ],
        )
        if chart:
            _add_picture(document, chart)
    else:
        document.add_paragraph("Noch keine Solarzeitreihe verfügbar.")

    document.add_heading("Gemessener vZEV-Energiefluss", level=1)
    flow_diagram = _flow_diagram(summary.get("zev_flow"))
    if flow_diagram:
        document.add_paragraph(
            "Die Grafik bildet den gemessenen CKW-ZEV-Aggregatfluss aus PV-Produktion, lokalem Verbrauch, "
            "Netzbezug und Rückspeisung ab."
        )
        _add_picture(document, flow_diagram)
    else:
        document.add_paragraph("Noch kein CKW-ZEV-Aggregat hochgeladen; deshalb liegt kein gemessener vZEV-Fluss vor.")

    document.add_heading("Lastverschiebung Nacht zu Tag", level=1)
    load_totals = load_shifting.get("totals", {})
    selected_profile = next(
        (
            profile
            for profile in load_shifting.get("profiles", [])
            if profile.get("file_id") == load_shifting.get("selected_consumption_file_id")
        ),
        None,
    )
    _add_key_values(
        document,
        [
            ("Kalenderjahr", str(load_shifting.get("selected_year") or "–")),
            ("Visualisierte Auswahl", str((selected_profile or {}).get("label") or "Alle Verbrauchszeitreihen")),
            ("Verbrauch Auswahl", _format_kwh(load_totals.get("annual_demand_kwh"))),
            ("Nachtverbrauch 22:00–06:00", _format_kwh(load_totals.get("nighttime_demand_kwh"))),
            ("Nachtanteil", _format_pct(load_totals.get("nighttime_share"))),
            ("Tagesverbrauch 08:00–18:00", _format_kwh(load_totals.get("daytime_demand_kwh"))),
            ("Anzahl Profile", _format_number(load_totals.get("profile_count"))),
        ],
    )
    heatmap = _load_shifting_heatmap(load_shifting.get("heatmap", {}))
    if heatmap:
        _add_picture(document, heatmap)
    if load_shifting.get("profiles"):
        _add_table(
            document,
            headers=["Profil", "Verbrauch", "Nacht", "Nachtanteil", "Ø/Nacht", "Tag", "Tage"],
            rows=[
                [
                    str(item["label"]),
                    _format_kwh(item["annual_demand_kwh"]),
                    _format_kwh(item["nighttime_demand_kwh"]),
                    _format_pct(item["nighttime_share"]),
                    _format_kwh(item["average_nightly_demand_kwh"]),
                    _format_kwh(item["daytime_demand_kwh"]),
                    _format_number(item["coverage_days"]),
                ]
                for item in load_shifting["profiles"]
            ],
        )
    else:
        document.add_paragraph("Noch kein Verbrauchsprofil für die Lastverschiebungsanalyse verfügbar.")

    document.add_heading("vZEV-Simulation", level=1)
    if simulation:
        sim_summary = simulation.summary
        sim_totals = sim_summary["totals"]
        pv_tariff = sim_summary["pv_tariff"]
        document.add_paragraph(
            "Die Simulation führt klassische PV- und Verbrauchszeitreihen je 15-Minuten-Intervall zu einem "
            "synthetischen vZEV zusammen und rechnet jeden Kalendermonat mit den CKW-Tarifen ab."
        )
        document.add_paragraph(
            (
                f"PV-Tarif: fixer Tarif von {_format_rp(pv_tariff['fixed_rp_per_kwh'])}; für jeden "
                "Teilnehmer und Monat gegen die 80%-Obergrenze geprüft. Überschreitungen werden als Warnung "
                "ausgewiesen, der gewählte Tarif wird dennoch verwendet."
                if pv_tariff["mode"] == "fixed"
                else "PV-Tarif: teilnehmerspezifisch 80% des vollständigen CKW-Standardproduktpreises."
            )
        )
        _add_table(
            document,
            headers=["Energieergebnis", "Menge", "Finanzergebnis", "Betrag"],
            rows=[
                ["Verbrauch", _format_kwh(sim_totals["consumption_kwh"]), "Kosten ohne vZEV", _format_chf(sim_totals["without_vzev_total_chf"])],
                ["PV-Produktion", _format_kwh(sim_totals["pv_production_kwh"]), "CKW-Kosten mit vZEV", _format_chf(sim_totals["with_vzev_ckw_invoice_chf"])],
                ["Lokale PV", _format_kwh(sim_totals["local_pv_consumption_kwh"]), "Kosten lokale PV", _format_chf(sim_totals["local_pv_energy_cost_chf"])],
                ["Netzbezug", _format_kwh(sim_totals["grid_import_kwh"]), "Kosten mit vZEV total", _format_chf(sim_totals["with_vzev_total_chf"])],
                ["Rückspeisung", _format_kwh(sim_totals["pv_excess_kwh"]), "Ersparnis", _format_chf(sim_totals["savings_chf"])],
                ["Lokale PV / Bedarf", _format_pct(sim_totals["self_consumption_pct"]), "Ersparnis Anteil", _format_pct(sim_totals["savings_pct"])],
                ["Lokale PV / PV-Produktion", _format_pct(sim_totals["pv_local_use_pct"]), "Tarifjahr", str(sim_summary["tariff_year"])],
            ],
        )
        energy_months = _simulation_monthly_energy(simulation)
        if energy_months:
            document.add_heading("Simulierte Energieflüsse nach Monat", level=2)
            _add_table(
                document,
                headers=["Monat", "Verbrauch", "PV-Produktion", "Lokale PV", "Netzbezug", "Rückspeisung"],
                rows=[
                    [
                        _month_label(item["month"]),
                        _format_kwh(item["consumption_kwh"]),
                        _format_kwh(item["pv_production_kwh"]),
                        _format_kwh(item["local_pv_kwh"]),
                        _format_kwh(item["grid_import_kwh"]),
                        _format_kwh(item["pv_excess_kwh"]),
                    ]
                    for item in energy_months
                ],
            )
            chart = _series_chart(
                energy_months,
                title="Simulierte vZEV-Energieflüsse",
                label_key="month",
                series=[
                    ("consumption_kwh", "Verbrauch", "#5167a8"),
                    ("pv_production_kwh", "PV-Produktion", "#d79c24"),
                    ("local_pv_kwh", "Lokale PV", "#287c71"),
                    ("grid_import_kwh", "Netzbezug", "#7b8a84"),
                    ("pv_excess_kwh", "Rückspeisung", "#9b4d55"),
                ],
            )
            if chart:
                _add_picture(document, chart)

        document.add_heading("Teilnehmer – Energie und Preise", level=2)
        _add_table(
            document,
            headers=[
                "Teilnehmer",
                "Verbrauch",
                "Netzbezug",
                "Lokale PV",
                "Standardprodukt",
                "80%-Obergrenze",
                "PV-Tarif",
                "Prüfung",
            ],
            rows=[
                [
                    str(item["participant"]),
                    _format_kwh(item["consumption_kwh"]),
                    _format_kwh(item["allocated_grid_import_kwh"]),
                    _format_kwh(item["allocated_local_pv_kwh"]),
                    _format_rp(item["standard_product_reference_rp_per_kwh"]),
                    _format_rp(item["local_energy_cap_rp_per_kwh"]),
                    _format_rp(item["local_energy_rp_per_kwh"]),
                    "Über 80%" if item["local_energy_tariff_exceeds_cap"] else "Innerhalb 80%",
                ]
                for item in sim_summary["participants"]
            ],
        )
        document.add_heading("Teilnehmer – Kostenvergleich", level=2)
        _add_table(
            document,
            headers=["Teilnehmer", "Ohne vZEV", "CKW-Anteil", "Lokale PV", "Mit vZEV total", "Ersparnis"],
            rows=[
                [
                    str(item["participant"]),
                    _format_chf(item["without_vzev_total_chf"]),
                    _format_chf(item["with_vzev_ckw_invoice_chf"]),
                    _format_chf(item["local_pv_energy_cost_chf"]),
                    _format_chf(item["with_vzev_total_chf"]),
                    _format_chf(item["savings_chf"]),
                ]
                for item in sim_summary["participants"]
            ],
        )
        document.add_heading("Tarif- und Leistungspreisänderung je Messpunkt", level=2)
        document.add_paragraph(
            "Der eigenständige Tarif jedes Messpunkts wird mit dem Tarif des aggregierten vZEV-Netzbezugs "
            "verglichen. Ein Wechsel über die Grenze von 50 MWh/Jahr kann insbesondere den Leistungspreis "
            "pro kW und Monat verändern."
        )
        _add_table(
            document,
            headers=["Teilnehmer / Messpunkt", "Jahresverbrauch einzeln", "Tarif einzeln", "vZEV-Jahresnetzbezug", "Tarif vZEV", "Leistungspreis einzeln", "Leistungspreis vZEV", "Tarifwechsel"],
            rows=[
                [
                    f"{item['participant']}\n{item.get('metering_code') or 'Messpunkt nicht angegeben'}",
                    _format_kwh(item["standalone_annualized_kwh"]),
                    _tariff_tier_label(item["standalone_tariff_tier"]),
                    _format_kwh(item["community_annualized_grid_import_kwh"]),
                    _tariff_tier_label(item["community_tariff_tier"]),
                    f"{_format_chf(item['standalone_demand_rate_chf_per_kw_month'])} / kW-Mt.",
                    f"{_format_chf(item['community_demand_rate_chf_per_kw_month'])} / kW-Mt.",
                    "Ja" if item["tariff_changed"] else "Nein",
                ]
                for item in sim_summary["participants"]
            ],
        )
        document.add_heading("Kostenänderung je Messpunkt", level=2)
        document.add_paragraph(
            "Vorzeichen: Plus bedeutet Mehrkosten, Minus bedeutet Kostenreduktion. Die CKW-Kostenänderung "
            "enthält keine interne PV-Bezahlung; die Gesamtkostenänderung enthält sie."
        )
        _add_table(
            document,
            headers=["Teilnehmer / Messpunkt", "Leistung einzeln", "Leistung vZEV-Anteil", "Änderung Leistung", "Änderung übrige CKW", "Änderung CKW total", "Lokale PV", "Änderung total"],
            rows=[
                [
                    f"{item['participant']}\n{item.get('metering_code') or 'Messpunkt nicht angegeben'}",
                    _format_chf(item["standalone_demand_cost_chf"]),
                    _format_chf(item["allocated_community_demand_cost_chf"]),
                    _format_cost_change(item["demand_cost_change_chf"]),
                    _format_cost_change(item["other_ckw_cost_change_chf"]),
                    _format_cost_change(item["ckw_cost_change_chf"]),
                    _format_chf(item["local_pv_energy_cost_chf"]),
                    _format_cost_change(item["total_cost_change_chf"]),
                ]
                for item in sim_summary["participants"]
            ],
        )
        _add_table(
            document,
            headers=["Gemeinschaft", "Änderung Leistung", "Änderung übrige CKW", "Änderung CKW total", "Lokale PV", "Änderung total"],
            rows=[
                [
                    "Total",
                    _format_cost_change(sim_totals["demand_cost_change_chf"]),
                    _format_cost_change(sim_totals["other_ckw_cost_change_chf"]),
                    _format_cost_change(sim_totals["ckw_cost_change_chf"]),
                    _format_chf(sim_totals["local_pv_energy_cost_chf"]),
                    _format_cost_change(sim_totals["total_cost_change_chf"]),
                ]
            ],
        )
        cost_change_chart = _vzev_cost_change_chart(sim_summary["participants"])
        if cost_change_chart:
            _add_picture(document, cost_change_chart)
        if sim_summary.get("monthly"):
            document.add_heading("Finanzergebnisse nach Monat", level=2)
            _add_table(
                document,
                headers=["Monat", "Ohne vZEV", "CKW", "Lokale PV", "Mit vZEV total", "Ersparnis"],
                rows=[
                    [
                        _month_label(item["month"]),
                        _format_chf(item["without_vzev_total_chf"]),
                        _format_chf(item["with_vzev_ckw_invoice_chf"]),
                        _format_chf(item["local_pv_energy_cost_chf"]),
                        _format_chf(item["with_vzev_total_chf"]),
                        _format_chf(item["savings_chf"]),
                    ]
                    for item in sim_summary["monthly"]
                ],
            )
    else:
        document.add_paragraph(f"Simulation nicht verfügbar: {simulation_error or 'Voraussetzungen fehlen.'}")

    document.add_heading("Batterieplatzierung", level=1)
    battery_input = battery.get("battery", {})
    _add_key_values(
        document,
        [
            ("Speicherkapazität", _format_kwh(battery_input.get("capacity_kwh"))),
            ("Maximale Ladeleistung", _format_kw(battery_input.get("max_charge_kw"))),
            ("Maximale Entladeleistung", _format_kw(battery_input.get("max_discharge_kw"))),
            ("Bestes PV-Szenario", str((battery.get("best_pv_site") or {}).get("label") or "kein Ergebnis")),
            ("Bestes Peak-Shaving-Szenario", str((battery.get("best_consumption_site") or {}).get("label") or "kein Ergebnis")),
        ],
    )
    document.add_heading("Fall 1: Batterie bei der PV-Produktion", level=2)
    if battery.get("pv_sites"):
        _add_table(
            document,
            headers=["PV-Datenstrom", "Zugeordnete Verbraucher", "PV-Energie", "Zusätzlich lokal", "Anteil", "Rückspeisung", "Speicherende"],
            rows=[
                [
                    str(item["label"]),
                    ", ".join(str(consumer["label"]) for consumer in item["assigned_consumers"]) if item["configured"] else "Zuordnung fehlt",
                    _format_kwh(item["pv_kwh"]),
                    _format_kwh(item["additional_local_solar_kwh"]) if item["configured"] else "–",
                    _format_pct(item["additional_local_solar_pct"]) if item["configured"] else "–",
                    _format_kwh(item["remaining_export_kwh"]) if item["configured"] else "–",
                    _format_kwh(item["ending_state_of_charge_kwh"]) if item["configured"] else "–",
                ]
                for item in battery["pv_sites"]
            ],
        )
        chart = _battery_chart(battery["pv_sites"], mode="pv")
        if chart:
            _add_picture(document, chart)
    else:
        document.add_paragraph("Noch keine PV-Zeitreihe für die Batteriesimulation verfügbar.")

    document.add_heading("Fall 2: Peak Shaving beim Verbraucher", level=2)
    if battery.get("consumption_sites"):
        _add_table(
            document,
            headers=["Verbrauchsdatenstrom", "Tarif", "Rate", "Spitzen vorher", "Mit Batterie", "Reduktion", "Einsparung"],
            rows=[
                [
                    str(item["label"]),
                    str(item["tariff_tier"]),
                    f"{_format_chf(item['demand_rate_chf_per_kw_month'])} / kW-Mt.",
                    _format_kw(item["baseline_peak_kw_month_sum"]),
                    _format_kw(item["battery_peak_kw_month_sum"]),
                    _format_kw(item["reduced_peak_kw_month_sum"]),
                    _format_chf(item["demand_charge_savings_chf"]),
                ]
                for item in battery["consumption_sites"]
            ],
        )
        chart = _battery_chart(battery["consumption_sites"], mode="consumption")
        if chart:
            _add_picture(document, chart)
    else:
        document.add_paragraph("Noch keine Verbrauchszeitreihe für die Peak-Shaving-Simulation verfügbar.")

    document.add_heading("Berechnungen und Annahmen", level=1)
    document.add_paragraph("vZEV-Energiefluss", style="Heading 2")
    _add_bullets(
        document,
        [
            "Lokale PV je 15-Minuten-Intervall = Minimum aus PV-Produktion und Gesamtverbrauch.",
            "Netzbezug = Maximum aus Gesamtverbrauch minus lokale PV, 0.",
            "Rückspeisung = Maximum aus PV-Produktion minus lokale PV, 0.",
            (
                f"Der interne PV-Tarif ist für alle Teilnehmer auf {_format_rp(fixed_local_energy_rp_per_kwh)} fixiert. Überschreitungen der jeweiligen 80%-Obergrenze werden gewarnt, aber in der Simulation nicht automatisch gekappt."
                if fixed_local_energy_rp_per_kwh is not None
                else "Der interne PV-Tarif entspricht je Teilnehmer 80% des vollständigen CKW-Standardproduktpreises; ausgewiesen werden sowohl Rp./kWh als auch CHF."
            ),
            "Kosten mit vZEV = zugeteilter CKW-Rechnungsanteil plus Kosten der lokal genutzten PV-Energie.",
            "Ersparnis = Kosten ohne vZEV minus Kosten mit vZEV.",
            "Leistungskostenänderung = zugeteilter Brutto-Leistungspreis des vZEV minus eigenständiger Brutto-Leistungspreis des Messpunkts.",
            "Übrige CKW-Kostenänderung = gesamte CKW-Kostenänderung minus Leistungskostenänderung.",
            "Bei Kostenänderungen bedeutet ein positiver Wert Mehrkosten und ein negativer Wert eine Kostenreduktion.",
        ],
    )
    assumptions = list(load_shifting.get("assumptions", [])) + list(battery.get("assumptions", []))
    if simulation:
        assumptions += list(simulation.summary.get("assumptions", []))
    _add_bullets(document, list(dict.fromkeys(str(value) for value in assumptions)))

    warnings = list(summary.get("warnings", [])) + list(plants.get("warnings", []))
    if simulation:
        warnings += list(simulation.summary.get("warnings", []))
    elif simulation_error:
        warnings.append(simulation_error)
    warnings = list(dict.fromkeys(str(value) for value in warnings))
    if warnings:
        document.add_heading("Hinweise und Datenqualität", level=1)
        _add_bullets(document, warnings)

    _style_tables(document)
    output = BytesIO()
    document.save(output)
    return WordReport(filename=f"vzev-simulationsbericht-{_slug(summary['name'])}.docx", content=output.getvalue())
